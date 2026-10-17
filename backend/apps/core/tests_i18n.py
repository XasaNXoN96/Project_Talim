import gettext as pygettext
import re
from pathlib import Path

from django.conf import settings
from django.contrib.auth.models import Group
from django.core import mail
from django.core.management import call_command
from django.test import Client, SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from django.utils import translation
from django.utils.translation import gettext

from apps.access.models import RoleAssignment
from apps.accounts.testing import PASSWORD, login_with_two_factor, make_user
from apps.documents.models import Document, DocumentCategory
from apps.documents.testing import PDF, upload
from apps.documents.services import add_version
from apps.documents.validation import validate_upload
from apps.notifications.models import Notification
from apps.organization.models import Department, Employee, Position
from apps.tasks.models import Task, TaskAttachment

LANGS = ("uz", "ru", "en")
LOCALE_DIR = Path(settings.BASE_DIR) / "locale"
PLACEHOLDER = re.compile(r"%\([A-Za-z_]+\)[sd]|%%|%s|%d")
UZ_APOSTROPHE = re.compile(r"[oOgG]‘")
CYRILLIC = re.compile(r"[А-Яа-яЁё]")


def parse_po(lang):
    text = (LOCALE_DIR / lang / "LC_MESSAGES" / "django.po").read_text()
    entries = []
    for block in re.split(r"\n\n+", text):
        mid = re.search(r"^msgid (.*?)\nmsgstr", block, flags=re.S | re.M)
        mstr = re.search(r"^msgstr (.*)$", block, flags=re.S | re.M)
        if not mid or not mstr:
            continue
        ctx = re.search(r'^msgctxt "(.*)"$', block, flags=re.M)
        entries.append({
            "id": "".join(re.findall(r'"(.*)"', mid.group(1))), "str": "".join(re.findall(r'"(.*)"', mstr.group(1))),
            "ctx": ctx.group(1) if ctx else None, "fuzzy": "#, fuzzy" in block,
        })
    return [e for e in entries if e["id"]]


class CatalogTests(SimpleTestCase):
    def test_every_language_has_a_catalog_and_a_compiled_file(self):
        for lang in LANGS:
            self.assertTrue((LOCALE_DIR / lang / "LC_MESSAGES" / "django.po").exists(), lang)
            self.assertTrue((LOCALE_DIR / lang / "LC_MESSAGES" / "django.mo").exists(), f"{lang}: run compilemessages")

    def test_no_empty_and_no_fuzzy_translations_anywhere(self):
        for lang in LANGS:
            entries = parse_po(lang)
            self.assertGreater(len(entries), 500, lang)
            self.assertEqual([e["id"] for e in entries if not e["str"].strip()], [], f"{lang}: empty translations")
            self.assertEqual([e["id"] for e in entries if e["fuzzy"]], [], f"{lang}: fuzzy translations are not shown by Django")

    def test_all_catalogs_have_the_same_messages(self):
        keys = {lang: {(e["ctx"], e["id"]) for e in parse_po(lang)} for lang in LANGS}
        self.assertEqual(keys["uz"], keys["ru"])
        self.assertEqual(keys["uz"], keys["en"])

    def test_placeholders_and_line_breaks_match_the_source_message(self):
        for lang in ("ru", "en"):
            for e in parse_po(lang):
                self.assertEqual(sorted(PLACEHOLDER.findall(e["id"])), sorted(PLACEHOLDER.findall(e["str"])), (lang, e["id"]))
                self.assertEqual(e["id"].count("\\n"), e["str"].count("\\n"), (lang, e["id"]))

    def test_translations_are_real_not_copies_of_the_uzbek_source(self):
        # A handful of terms legitimately look the same in two languages (PDF, Excel). Everything else must differ.
        same_ok = {"Excel", "PDF", "2FA"}
        same_ok_en = same_ok | {"Administrator", "e-mail", "Portal"}  # identical in Uzbek and English
        for lang in ("ru", "en"):
            ok = same_ok_en if lang == "en" else same_ok
            copies = [e["id"] for e in parse_po(lang) if e["str"] == e["id"] and e["id"] not in ok and CYRILLIC.search(e["id"]) is None]
            self.assertEqual(copies, [], f"{lang}: untranslated (copied) messages")

    def test_russian_has_cyrillic_and_english_has_none_and_neither_has_uzbek_apostrophes(self):
        for e in parse_po("ru"):
            if e["id"] not in {"Excel", "PDF"}:
                self.assertTrue(CYRILLIC.search(e["str"]), e["id"])
            self.assertIsNone(UZ_APOSTROPHE.search(e["str"]), e["id"])
        for e in parse_po("en"):
            self.assertIsNone(CYRILLIC.search(e["str"]), e["id"])
            self.assertIsNone(UZ_APOSTROPHE.search(e["str"]), e["id"])

    def test_compiled_files_match_the_po_files(self):
        for lang in LANGS:
            with open(LOCALE_DIR / lang / "LC_MESSAGES" / "django.mo", "rb") as fh:
                catalog = pygettext.GNUTranslations(fh)._catalog
            for e in parse_po(lang):
                key = e["id"].replace("\\n", "\n") if e["ctx"] is None else f'{e["ctx"]}\x04{e["id"]}'.replace("\\n", "\n")
                self.assertEqual(catalog.get(key), e["str"].replace("\\n", "\n"), (lang, e["id"]))

    def test_terminology_is_consistent(self):
        with translation.override("ru"):
            self.assertEqual(gettext("Xodim"), "Сотрудник")
            self.assertEqual(gettext("Hujjatlar"), "Документы")
            self.assertEqual(gettext("Topshiriqlar"), "Задания")
        with translation.override("en"):
            self.assertEqual(gettext("Xodim"), "Employee")
            self.assertEqual(gettext("Hujjatlar"), "Documents")
            self.assertEqual(gettext("Topshiriqlar"), "Tasks")

    def test_task_new_status_has_its_own_grammatical_form(self):
        from apps.tasks.models import Task as T

        with translation.override("ru"):
            self.assertEqual(str(T.Status.NEW.label), "Новое")
            self.assertEqual(gettext("Yangi"), "Новый")  # employee status
        with translation.override("en"):
            self.assertEqual(str(T.Status.NEW.label), "New")

    def test_django_own_messages_exist_for_all_three_languages(self):
        for lang in LANGS:
            with translation.override(lang):
                self.assertNotEqual(gettext("This field is required."), "")
        with translation.override("ru"):
            self.assertEqual(gettext("This field is required."), "Обязательное поле.")


class Base(TestCase):
    def setUp(self):
        call_command("seed_roles", verbosity=0)
        self.client = Client()

    def set_language(self, client, lang, next_url="/"):
        return client.post(reverse("set_language"), {"language": lang, "next": next_url})


class SwitchingTests(Base):
    def test_default_language_is_uzbek(self):
        r = self.client.get(reverse("accounts:login"))
        self.assertEqual(r["Content-Language"], "uz")
        self.assertContains(r, '<html lang="uz">')
        self.assertContains(r, "Tizimga kirish")

    def test_switching_sets_the_cookie_and_the_page_follows(self):
        r = self.set_language(self.client, "ru", reverse("accounts:login"))
        self.assertRedirects(r, reverse("accounts:login"), fetch_redirect_response=False)
        self.assertEqual(self.client.cookies[settings.LANGUAGE_COOKIE_NAME].value, "ru")
        page = self.client.get(reverse("accounts:login"))
        self.assertContains(page, '<html lang="ru">')
        self.assertContains(page, "Вход в систему")
        self.assertEqual(page["Content-Language"], "ru")

    def test_cookie_attributes(self):
        self.set_language(self.client, "en")
        c = self.client.cookies[settings.LANGUAGE_COOKIE_NAME]
        self.assertEqual((c["samesite"], bool(c["httponly"]), c["path"]), ("Lax", True, "/"))
        self.assertGreater(int(c["max-age"]), 60 * 60 * 24 * 30)

    def test_the_choice_is_kept_across_pages_and_anonymous_pages(self):
        self.set_language(self.client, "en")
        for url in (reverse("accounts:login"), reverse("content:home"), reverse("content:news_list"), reverse("accounts:password_reset")):
            page = self.client.get(url)
            self.assertContains(page, '<html lang="en">', msg_prefix=url)

    def test_current_page_and_filters_are_preserved(self):
        target = reverse("content:news_list") + "?page=2&q=maktab"
        r = self.set_language(self.client, "ru", target)
        self.assertEqual(r["Location"], target)
        page = self.client.get(reverse("content:news_list") + "?page=1")
        self.assertContains(page, 'name="next" value="/yangiliklar/?page=1"')

    def test_switcher_on_every_family_of_pages_and_marks_the_active_language(self):
        self.set_language(self.client, "ru")
        for url in (reverse("accounts:login"), reverse("content:home"), reverse("accounts:password_reset")):
            body = self.client.get(url).content.decode()
            self.assertIn(reverse("set_language"), body, url)
            self.assertRegex(body, r'<button[^>]*value="ru"[^>]*class="active"[^>]*aria-current="true"')
            for code in LANGS:
                self.assertRegex(body, rf'<button[^>]*value="{code}"[^>]*lang="{code}"')

    def test_open_redirect_is_refused(self):
        for evil in ("https://evil.test/x", "//evil.test", "javascript:alert(1)"):
            r = self.set_language(self.client, "ru", evil)
            self.assertEqual(r.status_code, 302)
            self.assertFalse(r["Location"].startswith(("http", "//", "javascript")), (evil, r["Location"]))

    def test_unknown_language_is_ignored(self):
        self.set_language(self.client, "ru")
        self.set_language(self.client, "de")
        self.set_language(self.client, "../etc")
        self.assertEqual(self.client.cookies[settings.LANGUAGE_COOKIE_NAME].value, "ru")

    def test_get_is_not_allowed_and_csrf_is_enforced(self):
        self.assertEqual(self.client.get(reverse("set_language")).status_code, 405)
        c = Client(enforce_csrf_checks=True)
        self.assertEqual(c.post(reverse("set_language"), {"language": "ru", "next": "/"}).status_code, 403)

    def test_accept_language_is_used_only_without_a_cookie(self):
        self.assertContains(self.client.get(reverse("accounts:login"), headers={"accept-language": "ru-RU,ru;q=0.9"}), '<html lang="ru">')
        self.assertContains(self.client.get(reverse("accounts:login"), headers={"accept-language": "fr"}), '<html lang="uz">')
        self.set_language(self.client, "en")
        self.assertContains(self.client.get(reverse("accounts:login"), headers={"accept-language": "ru"}), '<html lang="en">')

    def test_no_per_language_urls_exist(self):
        for prefix in ("/ru/", "/en/", "/uz/"):
            self.assertEqual(self.client.get(prefix).status_code, 404, prefix)


class SignedInTests(Base):
    def setUp(self):
        super().setUp()
        self.user = make_user("alice", role="employee")

    def test_choice_is_stored_on_the_account_and_follows_to_another_browser(self):
        c1 = Client()
        login_with_two_factor(c1, self.user)
        self.set_language(c1, "ru", reverse("portal:dashboard"))
        self.user.refresh_from_db()
        self.assertEqual(self.user.locale, "ru")
        c2 = Client()  # a different browser: no cookie
        login_with_two_factor(c2, self.user)
        self.assertContains(c2.get(reverse("portal:dashboard")), '<html lang="ru">')

    def test_cookie_wins_over_the_stored_language(self):
        self.user.locale = "ru"
        self.user.save()
        c = Client()
        login_with_two_factor(c, self.user)
        self.set_language(c, "en")
        self.assertContains(c.get(reverse("portal:dashboard")), '<html lang="en">')

    def test_a_signed_out_switch_does_not_touch_any_account(self):
        self.set_language(self.client, "en")
        self.user.refresh_from_db()
        self.assertEqual(self.user.locale, "uz")

    def test_switching_works_for_users_forced_to_change_password_or_enrol_2fa(self):
        self.user.must_change_password = True
        self.user.save()
        c = Client()
        login_with_two_factor(c, self.user)
        r = self.set_language(c, "ru", "/")
        self.assertEqual(r.status_code, 302)
        self.assertNotIn("/accounts/password-change/", r["Location"])  # the endpoint is reachable, not bounced to the change form


class DataFixtureMixin:
    """A small organisation with one object of every kind, used to render pages in every language."""

    def build(self):
        call_command("seed_roles", verbosity=0)
        self.dept, self.sub = Department.objects.create(name="Metodika"), None
        self.sub = Department.objects.create(name="Maktabgacha", parent=self.dept)
        self.pos = Position.objects.create(name="Metodist")
        self.admin = make_user("boss", role="admin")
        self.hr = make_user("hrista", role="hr")
        self.emp = make_user("plain", role="employee")
        self.card = Employee.objects.create(last_name="Aliyev", first_name="Vali", department=self.sub, position=self.pos, user=self.emp, status="active", work_phone="+998 71 111 22 33")
        self.task = Task.objects.create(title="Hisobot", created_by=self.admin, status="review")
        self.task.assignees.set([self.emp])
        self.category = DocumentCategory.objects.get(name="Buyruqlar")
        self.doc = Document.objects.create(title="Nizom", category=self.category, created_by=self.hr)
        add_version(self.doc, validate_upload(upload("a.pdf")), actor=self.hr)
        self.attachment = None
        from apps.content.models import News, NewsCategory

        self.news_cat = NewsCategory.objects.create(name="Maktab", slug="maktab")
        self.news = News.objects.create(title="Maktab xabari", slug="maktab-xabari", body="Matn", category=self.news_cat, status="draft")
        self.news_review = News.objects.create(title="Tekshiruvdagi xabar", slug="tekshiruvda", body="Matn", category=self.news_cat, status="review")
        Notification.objects.create(user=self.emp, kind="task_assigned", task=self.task, title="Hisobot")

    def client_for(self, user, lang):
        c = Client()
        login_with_two_factor(c, user)
        c.cookies[settings.LANGUAGE_COOKIE_NAME] = lang
        return c


def visible_text(html):
    """The page without the language switcher and the language <select> options: both name each language in that
    language on purpose (O‘zbekcha / Русский / English)."""
    html = re.sub(r'<form class="lang-switch".*?</form>', "", html, flags=re.S)
    return re.sub(r"<option[^>]*>(O‘zbekcha|Русский|English)</option>", "", html)


class EveryPageInEveryLanguage(DataFixtureMixin, TestCase):
    def setUp(self):
        self.build()

    def pages(self):
        t, d, e = self.task.pk, self.doc.pk, self.card.pk
        return [
            (None, "accounts:login", ()), (None, "accounts:password_reset", ()), (None, "content:home", ()), (None, "content:news_list", ()),
            ("admin", "portal:dashboard", ()), ("admin", "portal:audit", ()), ("admin", "accounts:security", ()), ("admin", "accounts:password_change", ()),
            ("admin", "users:list", ()), ("admin", "users:create", ()), ("admin", "users:detail", (self.emp.pk,)), ("admin", "users:edit", (self.emp.pk,)),
            ("admin", "employees:list", ()), ("admin", "employees:detail", (e,)), ("admin", "employees:create", ()), ("admin", "employees:edit", (e,)),
            ("admin", "employees:structure", ()), ("admin", "employees:department_create", ()), ("admin", "employees:position_create", ()),
            ("admin", "tasks:list", ()), ("admin", "tasks:detail", (t,)), ("admin", "tasks:create", ()), ("admin", "tasks:edit", (t,)),
            ("admin", "documents:list", ()), ("admin", "documents:detail", (d,)), ("admin", "documents:create", ()), ("admin", "documents:edit", (d,)),
            ("emp", "notifications:list", ()), ("emp", "tasks:detail", (t,)), ("emp", "portal:dashboard", ()),
            ("admin", "reports:index", ()),
            ("admin", "documents:preview", (d,)), ("admin", "news_admin:list", ()), ("admin", "news_admin:create", ()), ("admin", "news_admin:preview", (self.news.pk,)),
            ("admin", "news_admin:edit", (self.news.pk,)), ("admin", "news_admin:categories", ()), ("admin", "news_admin:category_create", ()),
            ("admin", "news_admin:category_edit", (self.news_cat.pk,)), ("admin", "news_admin:status", (self.news.pk, "submit")),
            ("admin", "news_admin:status", (self.news_review.pk, "approve")), ("admin", "news_admin:status", (self.news_review.pk, "request_changes")), ("admin", "news_admin:preview", (self.news_review.pk,)),
            ("admin", "employees:import", ()), ("admin", "portal:search", ()), ("emp", "portal:search", ()),
            ("admin", "portal:jobs", ()), ("admin", "accounts:sessions", ()), ("admin", "accounts:sessions_end_others", ()), ("emp", "accounts:sessions", ()),
        ]

    def test_pages_render_in_each_language_with_the_right_script(self):
        users = {"admin": self.admin, "emp": self.emp, None: None}
        for lang in LANGS:
            for who, name, args in self.pages():
                c = self.client_for(users[who], lang) if who else Client()
                if not who:
                    c.cookies[settings.LANGUAGE_COOKIE_NAME] = lang
                url = reverse(name, args=args)
                r = c.get(url)
                self.assertEqual(r.status_code, 200, (lang, url))
                body = r.content.decode()
                self.assertIn(f'<html lang="{lang}">', body, (lang, url))
                text = visible_text(body)
                self.assertIsNone(UZ_APOSTROPHE.search(text) if lang != "uz" else None, (lang, url, UZ_APOSTROPHE.findall(text)[:3], re.findall(r".{20}[oOgG]‘.{10}", text)[:2]))
                if lang == "ru":
                    self.assertTrue(CYRILLIC.search(text), (lang, url))
                if lang == "en":
                    self.assertIsNone(CYRILLIC.search(text), (lang, url, re.findall(r".{15}[А-Яа-я]+.{10}", text)[:2]))

    def test_pages_in_russian_and_english_do_not_contain_uzbek_interface_words(self):
        uz_words = ("Bosh sahifa", "Saqlash", "Bekor qilish", "Qidirish", "Xodimlar", "Topshiriqlar", "Hujjatlar", "Foydalanuvchilar", "Yangilik", "Tizimdan chiqish")
        for lang in ("ru", "en"):
            for who, name, args in self.pages():
                users = {"admin": self.admin, "emp": self.emp}
                c = self.client_for(users[who], lang) if who else Client()
                if not who:
                    c.cookies[settings.LANGUAGE_COOKIE_NAME] = lang
                body = visible_text(c.get(reverse(name, args=args)).content.decode())
                # fixture data (names typed by users, e.g. the task called "Hisobot") is allowed; interface words are not
                for w in uz_words:
                    self.assertNotIn(w, re.sub(r"Hisobot|Nizom|Maktab xabari|Maktab|Metodika|Maktabgacha|Aliyev|Metodist|Vali", "", body), (lang, name, w))

    def test_error_pages_are_translated(self):
        for lang, expected in (("uz", "Sahifa topilmadi"), ("ru", "Страница не найдена"), ("en", "Page not found")):
            c = Client()
            c.cookies[settings.LANGUAGE_COOKIE_NAME] = lang
            r = c.get("/no-such-page/")
            self.assertEqual(r.status_code, 404)
            self.assertContains(r, expected, status_code=404)
            self.assertContains(r, f'<html lang="{lang}">', status_code=404)
        for lang, expected in (("ru", "Доступ запрещён"), ("en", "Access denied"), ("uz", "Ruxsat yo‘q")):
            c = self.client_for(self.emp, lang)
            r = c.get(reverse("users:list"))
            self.assertContains(r, expected, status_code=403)

    def test_500_page_renders_translated_without_a_request(self):
        from apps.core.views import error_500

        for lang, expected in (("ru", "Ошибка сервера"), ("en", "Server error"), ("uz", "Serverda xatolik")):
            with translation.override(lang):
                r = error_500(None)
            self.assertEqual(r.status_code, 500)
            self.assertIn(expected, r.content.decode())

    def test_csrf_failure_page_is_translated(self):
        c = Client(enforce_csrf_checks=True)
        c.cookies[settings.LANGUAGE_COOKIE_NAME] = "ru"
        r = c.post(reverse("accounts:login"), {"username": "a", "password": "b"})
        self.assertContains(r, "Проверка безопасности не пройдена", status_code=403)


class FlashMessagesAndValidationTests(DataFixtureMixin, TestCase):
    def setUp(self):
        self.build()

    def test_success_messages_follow_the_language(self):
        expected = {"uz": "Foydalanuvchi yaratildi", "ru": "Пользователь создан", "en": "The user was created"}
        for i, lang in enumerate(LANGS):
            c = self.client_for(self.admin, lang)
            data = {"username": f"new{i}", "email": f"new{i}@example.test", "last_name": "A", "first_name": "B", "role": "employee"}
            r = c.post(reverse("users:create"), data, follow=True)
            self.assertContains(r, expected[lang])

    def test_validation_errors_follow_the_language(self):
        expected = {"uz": "Bu login band.", "ru": "Этот логин занят.", "en": "This username is already taken."}
        for lang in LANGS:
            c = self.client_for(self.admin, lang)
            r = c.post(reverse("users:create"), {"username": "plain", "email": "x@example.test", "last_name": "A", "first_name": "B", "role": "employee"})
            self.assertContains(r, expected[lang])

    def test_django_builtin_form_errors_follow_the_language(self):
        for lang, text in (("ru", "Обязательное поле."), ("en", "This field is required."), ("uz", None)):
            c = self.client_for(self.admin, lang)
            r = c.post(reverse("users:create"), {})
            if text:
                self.assertContains(r, text)

    def test_login_error_and_lockout_message_are_translated_and_identical_across_causes(self):
        for lang, expected in (("ru", "Неверный логин или пароль"), ("en", "The login or password is incorrect")):
            msgs = []
            for username in ("plain", "no-such-user"):
                c = Client()
                c.cookies[settings.LANGUAGE_COOKIE_NAME] = lang
                r = c.post(reverse("accounts:login"), {"username": username, "password": "wrong-password"})
                self.assertContains(r, expected)
                msgs.append(re.search(r'role="alert">.*?<span>(.*?)</span>', r.content.decode(), flags=re.S).group(1))
            self.assertEqual(msgs[0], msgs[1])  # still no account enumeration

    def test_upload_rejections_are_translated(self):
        for lang, text in (("ru", "Размер файла не должен превышать"), ("en", "The file size must not exceed")):
            c = self.client_for(self.hr, lang)
            with override_settings(DOCUMENT_MAX_BYTES=50):
                r = c.post(reverse("documents:create"), {"title": "T", "category": self.category.pk, "access_level": "general", "file": upload("x.pdf", PDF + b"0" * 100)})
            self.assertContains(r, text)
        c = self.client_for(self.hr, "ru")
        r = c.post(reverse("documents:create"), {"title": "T", "category": self.category.pk, "access_level": "general", "file": upload("x.exe", b"MZ")})
        self.assertContains(r, "Этот тип файла не разрешён")

    def test_report_filter_errors_are_translated(self):
        for lang, text in (("ru", "Дата начала не может быть позже даты окончания."), ("en", "The start date cannot be later than the end date.")):
            c = self.client_for(self.admin, lang)
            r = c.get(reverse("reports:index"), {"date_from": "2030-01-02", "date_to": "2030-01-01"})
            self.assertContains(r, text)


class LocalisedOutputTests(DataFixtureMixin, TestCase):
    def setUp(self):
        self.build()

    def test_users_and_roles_are_shown_with_translated_names_not_codes(self):
        c = self.client_for(self.admin, "ru")
        body = c.get(reverse("users:detail", args=[self.hr.pk])).content.decode()
        self.assertIn("Отдел кадров", body)
        self.assertNotIn(">hr<", body)
        body = self.client_for(self.admin, "en").get(reverse("users:list")).content.decode()
        self.assertIn("HR department", body)
        self.assertIn("Administrator", body)

    def test_task_history_shows_status_names_not_codes(self):
        w = self.client_for(self.admin, "ru")
        w.post(reverse("tasks:transition", args=[self.task.pk, "done"]))
        body = w.get(reverse("tasks:detail", args=[self.task.pk])).content.decode()
        self.assertIn("На проверке → Выполнено", body)
        self.assertNotIn("review →", body)
        body_en = self.client_for(self.admin, "en").get(reverse("tasks:detail", args=[self.task.pk])).content.decode()
        self.assertIn("In review → Done", body_en)

    def test_status_and_priority_labels(self):
        t = Task.objects.create(title="X", created_by=self.admin, status="new", priority="high")
        t.assignees.set([self.emp])
        for lang, new, high in (("uz", "Yangi", "Yuqori"), ("ru", "Новое", "Высокий"), ("en", "New", "High")):
            body = self.client_for(self.admin, lang).get(reverse("tasks:list"), {"tab": "all"}).content.decode()
            self.assertIn(new, body)
            self.assertIn(high, body)

    def test_notification_kinds_are_translated(self):
        for lang, text in (("ru", "Вам назначено задание"), ("en", "A task was assigned to you")):
            self.assertContains(self.client_for(self.emp, lang).get(reverse("notifications:list")), text)

    def test_audit_log_actions_are_translated_but_stored_codes_are_unchanged(self):
        c = self.client_for(self.admin, "en")
        c.post(reverse("tasks:transition", args=[self.task.pk, "done"]))
        from apps.audit.models import AuditLog

        AuditLog.objects.create(action=AuditLog.Action.DOCUMENT_DOWNLOADED, user=self.admin)
        entry = AuditLog.objects.latest("at")
        self.assertEqual(entry.action, "document_downloaded")  # database value never depends on the language
        self.assertContains(c.get(reverse("portal:audit")), "Document downloaded")
        self.assertContains(self.client_for(self.admin, "ru").get(reverse("portal:audit")), "Документ скачан")

    def test_numbers_dates_and_dynamic_values_are_not_broken(self):
        body = self.client_for(self.admin, "ru").get(reverse("documents:detail", args=[self.doc.pk])).content.decode()
        self.assertRegex(body, r"\d{2}\.\d{2}\.\d{4} \d{2}:\d{2}")  # dd.mm.yyyy hh:mm in every language
        self.assertNotIn("%(", body)
        self.assertNotIn("{{", body)
        self.assertNotIn("%s", body)

    def test_reports_page_and_excel_export_in_each_language(self):
        from io import BytesIO

        from openpyxl import load_workbook

        sheets = {"uz": "Umumiy", "ru": "Общее", "en": "General"}
        kpi = {"uz": "Faol xodimlar", "ru": "Активные сотрудники", "en": "Active employees"}
        for lang in LANGS:
            c = self.client_for(self.admin, lang)
            page = c.get(reverse("reports:index"))
            self.assertContains(page, kpi[lang])
            x = c.get(reverse("reports:export", args=["xlsx"]))
            wb = load_workbook(BytesIO(x.content))
            self.assertEqual(wb.sheetnames[0], sheets[lang])
            values = {r[0]: r[1] for r in wb.worksheets[0].iter_rows(values_only=True) if len(r) > 1 and r[1] is not None}
            self.assertEqual(values[kpi[lang]], 1)  # the same number in every language
            pdf = c.get(reverse("reports:export", args=["pdf"]))
            self.assertTrue(pdf.content.startswith(b"%PDF-"))

    def test_report_pdf_and_excel_filenames_do_not_depend_on_language(self):
        for lang in LANGS:
            r = self.client_for(self.admin, lang).get(reverse("reports:export", args=["xlsx"]))
            self.assertRegex(r["Content-Disposition"], r'filename="hisobot-\d{8}-\d{4}\.xlsx"')


class EmailLanguageTests(DataFixtureMixin, TestCase):
    def setUp(self):
        self.build()

    def test_invitation_is_written_in_the_recipients_language_not_the_senders(self):
        self.emp.locale = "ru"
        self.emp.save()
        c = self.client_for(self.admin, "en")  # the administrator works in English
        c.post(reverse("users:action", args=[self.emp.pk, "invite"]))
        m = mail.outbox[-1]
        self.assertIn("Здравствуйте", m.body)
        self.assertIn("приглашение", m.subject)
        self.assertNotIn("Hello", m.body)

    def test_invitation_in_english_and_uzbek(self):
        for lang, hello in (("en", "Hello,"), ("uz", "Assalomu alaykum")):
            self.emp.locale = lang
            self.emp.save()
            from apps.accounts.throttle import clear

            clear([])
            self.client_for(self.admin, "ru").post(reverse("users:action", args=[self.emp.pk, "invite"]))
            self.assertIn(hello, mail.outbox[-1].body, lang)

    def test_security_notice_uses_the_account_language(self):
        from apps.accounts.notifications import notify_security_event

        for lang, text in (("ru", "Ваш пароль изменён"), ("en", "Your password was changed"), ("uz", "Parolingiz almashtirildi")):
            self.emp.locale = lang
            self.emp.save()
            notify_security_event(self.emp, "password_changed")
            self.assertIn(text, mail.outbox[-1].subject, lang)

    def test_password_reset_mail_language_and_link(self):
        self.emp.locale = "en"
        self.emp.save()
        c = Client()
        c.cookies[settings.LANGUAGE_COOKIE_NAME] = "ru"  # the requester browses in Russian; the mail follows the ACCOUNT
        c.post(reverse("accounts:password_reset"), {"email": self.emp.email})
        m = mail.outbox[-1]
        self.assertIn("password reset", m.subject)
        self.assertIn("To set a new password", m.body)
        self.assertIn("/accounts/reset/", m.body)

    @override_settings(NOTIFICATION_EMAILS_ENABLED=True)
    def test_notification_copy_uses_the_recipients_language(self):
        from apps.notifications.services import notify

        self.emp.locale = "ru"
        self.emp.save()
        with self.captureOnCommitCallbacks(execute=True):
            notify(self.emp, "task_overdue", self.task, dedupe_key="i18n-1")
        self.assertIn("Срок задания истёк", mail.outbox[-1].body)

    def test_emails_contain_no_untranslated_placeholders(self):
        for lang in LANGS:
            self.emp.locale = lang
            self.emp.save()
            from apps.accounts.notifications import notify_security_event

            for event in ("password_changed", "two_factor_enabled", "two_factor_disabled", "backup_codes_regenerated", "two_factor_reset", "password_reset", "invitation"):
                notify_security_event(self.emp, event)
                body = (mail.outbox[-1].subject + mail.outbox[-1].body).replace(settings.TWO_FACTOR_ISSUER, "")  # brand name stays as is
                self.assertNotIn("%(", body)
                if lang != "uz":
                    self.assertIsNone(UZ_APOSTROPHE.search(body), (lang, event))
                if lang == "en":
                    self.assertIsNone(CYRILLIC.search(body), (lang, event))


class AccessIsLanguageIndependentTests(DataFixtureMixin, TestCase):
    def setUp(self):
        self.build()

    def test_permissions_are_identical_in_every_language(self):
        outcomes = {}
        checks = [("emp", "users:list"), ("emp", "reports:index"), ("emp", "documents:create"), ("emp", "portal:audit"), ("admin", "reports:index"), ("hr", "users:list"), (None, "portal:dashboard")]
        users = {"admin": self.admin, "emp": self.emp, "hr": self.hr, None: None}
        for lang in LANGS:
            for who, name in checks:
                c = self.client_for(users[who], lang) if who else Client()
                if not who:
                    c.cookies[settings.LANGUAGE_COOKIE_NAME] = lang
                outcomes.setdefault((who, name), set()).add(c.get(reverse(name)).status_code)
        for key, codes in outcomes.items():
            self.assertEqual(len(codes), 1, f"{key} differs between languages: {codes}")

    def test_forbidden_pages_reveal_nothing_in_any_language(self):
        for lang in LANGS:
            c = self.client_for(self.emp, lang)
            r = c.get(reverse("users:detail", args=[self.admin.pk]))
            self.assertEqual(r.status_code, 403)
            self.assertNotIn(self.admin.email, r.content.decode())
            r = c.get(reverse("reports:index"))
            self.assertEqual(r.status_code, 403)
            self.assertNotIn("Aliyev", r.content.decode())

    def test_language_cookie_cannot_be_used_to_reach_other_data_or_scopes(self):
        c = self.client_for(self.emp, "ru")
        c.cookies[settings.LANGUAGE_COOKIE_NAME] = "../../etc/passwd"  # garbage in the cookie falls back to the default
        r = c.get(reverse("portal:dashboard"))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, '<html lang="uz">')


class FlowPagesInEveryLanguage(DataFixtureMixin, TestCase):
    """The pages that only exist in the middle of a flow (2FA enrolment, password reset, backup codes, ...)."""

    def setUp(self):
        self.build()

    def check(self, lang, response, label, expect_status=200):
        self.assertEqual(response.status_code, expect_status, (lang, label))
        body = response.content.decode()
        text = visible_text(body)
        self.assertIn(f'<html lang="{lang}">', body, (lang, label))
        if lang != "uz":
            self.assertIsNone(UZ_APOSTROPHE.search(text), (lang, label, re.findall(r".{20}[oOgG]‘.{10}", text)[:2]))
        if lang == "en":
            self.assertIsNone(CYRILLIC.search(text), (lang, label, re.findall(r".{15}[А-Яа-я]+.{10}", text)[:2]))
        if lang == "ru":
            self.assertTrue(CYRILLIC.search(text), (lang, label))
        return text

    def test_two_factor_pages(self):
        from apps.accounts.tests_twofactor import code_at, enrol, password_step

        for lang in LANGS:
            boss = make_user(f"boss_{lang}", role="admin")
            c = Client()
            c.cookies[settings.LANGUAGE_COOKIE_NAME] = lang
            password_step(c, boss.username)
            self.check(lang, c.get(reverse("accounts:two_factor_setup")), "2fa setup")
            secret, _codes = enrol(boss)
            c2 = Client()
            c2.cookies[settings.LANGUAGE_COOKIE_NAME] = lang
            password_step(c2, boss.username)
            self.check(lang, c2.get(reverse("accounts:two_factor_verify")), "2fa verify")
            wrong = c2.post(reverse("accounts:two_factor_verify"), {"code": "000000"})
            self.check(lang, wrong, "2fa verify wrong code")
            self.assertRegex(wrong.content.decode(), r'role="alert"')

    def test_security_backup_codes_and_disable_pages(self):
        from apps.accounts.tests_twofactor import enrol

        for lang in LANGS:
            u = make_user(f"hr_{lang}", role="hr")
            enrol(u)
            c = self.client_for(u, lang)
            self.check(lang, c.get(reverse("accounts:security")), "security (2fa on)")
            self.check(lang, c.get(reverse("accounts:backup_codes_regenerate")), "regenerate form")
            self.check(lang, c.get(reverse("accounts:two_factor_disable")), "disable form")
            done = c.post(reverse("accounts:backup_codes_regenerate"), {"password": PASSWORD}, follow=True)
            self.check(lang, done, "new backup codes")

    def test_password_reset_flow_pages(self):
        from django.contrib.auth.tokens import default_token_generator
        from django.utils.encoding import force_bytes
        from django.utils.http import urlsafe_base64_encode

        for lang in LANGS:
            c = Client()
            c.cookies[settings.LANGUAGE_COOKIE_NAME] = lang
            self.check(lang, c.get(reverse("accounts:password_reset")), "reset form")
            r = c.post(reverse("accounts:password_reset"), {"email": "plain@example.test"}, follow=True)
            self.check(lang, r, "reset done")
            uid, token = urlsafe_base64_encode(force_bytes(self.emp.pk)), default_token_generator.make_token(self.emp)
            link = reverse("accounts:password_reset_confirm", args=[uid, token])
            page = c.get(link, follow=True)
            self.check(lang, page, "reset confirm form")
            bad = c.get(reverse("accounts:password_reset_confirm", args=[uid, "bad-token"]))
            self.check(lang, bad, "reset invalid link")
            self.check(lang, c.get(reverse("accounts:password_reset_complete")), "reset complete")

    def test_password_change_pages_and_errors(self):
        for lang in LANGS:
            c = self.client_for(self.emp, lang)
            self.check(lang, c.get(reverse("accounts:password_change")), "change form")
            bad = c.post(reverse("accounts:password_change"), {"old_password": "wrong", "new_password1": "x", "new_password2": "y"})
            self.check(lang, bad, "change errors")
            self.check(lang, c.get(reverse("accounts:password_change_done")), "change done")

    def test_forced_password_change_page_and_language_switch_still_work(self):
        self.emp.must_change_password = True
        self.emp.save()
        for lang in LANGS:
            c = self.client_for(self.emp, lang)
            r = c.get(reverse("portal:dashboard"), follow=True)
            text = self.check(lang, r, "forced change")
            self.assertIn("/accounts/password-change/", r.redirect_chain[-1][0])

    def test_news_detail_and_empty_states(self):
        from django.utils import timezone

        from apps.content.models import News

        news = News.objects.create(title="Maktab yangiligi", slug="maktab", body="Matn", excerpt="Qisqa", status="published", publish_at=timezone.now())
        for lang in LANGS:
            c = Client()
            c.cookies[settings.LANGUAGE_COOKIE_NAME] = lang
            self.check(lang, c.get(reverse("content:news_detail", args=[news.slug])), "news detail")
            self.check(lang, c.get(reverse("content:news_list")), "news list")
        News.objects.all().delete()
        for lang in LANGS:
            c = Client()
            c.cookies[settings.LANGUAGE_COOKIE_NAME] = lang
            self.check(lang, c.get(reverse("content:news_list")), "news list empty")
            self.check(lang, c.get(reverse("content:home")), "home no news")

    def test_empty_states_in_portal(self):
        Task.objects.all().delete()
        Notification.objects.all().delete()
        Document.objects.all().delete()
        for lang in LANGS:
            c = self.client_for(self.admin, lang)
            for name in ("tasks:list", "notifications:list", "documents:list", "employees:list"):
                self.check(lang, c.get(reverse(name)), f"empty {name}")
            self.check(lang, c.get(reverse("reports:index"), {"date_from": "2001-01-01", "date_to": "2001-01-31"}), "empty report")

    def test_attachment_and_archived_states(self):
        for lang in LANGS:
            c = self.client_for(self.admin, lang)
            c.post(reverse("tasks:attachment_add", args=[self.task.pk]), {"file": upload("a.pdf", PDF + lang.encode())})
            self.check(lang, c.get(reverse("tasks:detail", args=[self.task.pk])), "task with attachment")
        self.doc.status = "archived"
        self.doc.save()
        for lang in LANGS:
            self.check(lang, self.client_for(self.admin, lang).get(reverse("documents:detail", args=[self.doc.pk])), "archived doc")


class DjangoAdminInEveryLanguage(DataFixtureMixin, TestCase):
    def setUp(self):
        self.build()
        from django.contrib.auth import get_user_model

        self.root = get_user_model().objects.create_superuser("root", "root@example.test", PASSWORD)

    def test_admin_pages_follow_the_language_and_our_models_are_translated(self):
        pages = ["", "accounts/user/", f"accounts/user/{self.emp.pk}/change/", "organization/employee/", "organization/employee/add/",
                 "audit/auditlog/", "documents/documentcategory/", "organization/department/", "access/roleassignment/", "content/news/"]
        heading = {"uz": None, "ru": "Администрирование Django", "en": "Django administration"}
        for lang in LANGS:
            c = self.client_for(self.root, lang)
            for page in pages:
                r = c.get("/" + settings.ADMIN_URL + page)
                self.assertEqual(r.status_code, 200, (lang, page))
                body = visible_text(r.content.decode())
                self.assertIn(f'lang="{lang}"', body, (lang, page))
                if lang != "uz":
                    self.assertIsNone(UZ_APOSTROPHE.search(re.sub(r"Mirzacho‘l XTB", "", body)), (lang, page, re.findall(r".{25}[oOgG]‘.{12}", body)[:2]))
                if lang == "en":
                    self.assertIsNone(CYRILLIC.search(body), (lang, page, re.findall(r".{15}[А-Яа-я]+.{10}", body)[:2]))
            index = c.get("/" + settings.ADMIN_URL).content.decode()
            if heading[lang]:
                self.assertIn(heading[lang], index)
        ru = self.client_for(self.root, "ru").get("/" + settings.ADMIN_URL).content.decode()
        for expected in ("Сотрудники", "Документы", "Назначения ролей", "Подразделения", "Пользователи"):
            self.assertIn(expected, ru)
        en = self.client_for(self.root, "en").get("/" + settings.ADMIN_URL + "organization/employee/add/").content.decode()
        for expected in ("last name", "first name", "hire date"):
            self.assertIn(expected, en.lower())  # the admin capitalises labels
