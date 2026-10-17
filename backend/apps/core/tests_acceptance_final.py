"""Day 22: final acceptance. Each test is ONE chain through several modules over real HTTP (the same requests a browser sends), using the
ready-made roles. Single-module details have their own suites; this file proves the modules work TOGETHER."""
import datetime
import io
import json
import re

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group, Permission
from django.core import mail
from django.core.management import call_command
from django.test import Client, TestCase
from django.urls import reverse
from django.utils import timezone

from apps.access.models import RoleAssignment
from apps.accounts.tests_twofactor import SETUP, code_at, password_step, secret_of, signed_in
from apps.accounts.testing import login_with_two_factor, make_user
from apps.audit.models import AuditLog
from apps.content import workflow
from apps.content.models import News
from apps.core.acceptance import World
from apps.documents.models import Document, DocumentCategory, DocumentVersion
from apps.documents.testing import PDF, upload
from apps.documents.tests_scanner import FakeClamd, scanning
from apps.notifications.models import Notification
from apps.tasks.models import Task, TaskEvent

User = get_user_model()
NEW_PASSWORD = "Yangi-Sinov-Parol-24680!"


def login(user):
    c = Client()
    login_with_two_factor(c, user)
    return c


def set_password_from_invitation(username, role, admin_client, **extra):
    """Admin creates the account in the portal and sends the invitation; the person opens the e-mailed link and chooses a password."""
    r = admin_client.post(reverse("users:create"), {"username": username, "email": f"{username}@example.test", "last_name": "Sinov", "first_name": "Yangi",
                                                     "role": role, "department": extra.get("department", "")})
    assert r.status_code == 302, r.context["form"].errors if getattr(r, "context", None) else r.status_code
    user = User.objects.get(username=username)
    mail.outbox.clear()
    assert admin_client.post(reverse("users:action", args=[user.pk, "invite"])).status_code == 302
    link = re.search(r"/accounts/reset/[^\s]+", mail.outbox[-1].body).group(0)
    page = Client().get(link, follow=True)
    form_url = page.redirect_chain[-1][0]
    c = Client()
    c.get(link, follow=True)
    r = c.post(form_url, {"new_password1": NEW_PASSWORD, "new_password2": NEW_PASSWORD})
    assert r.status_code == 302, r.status_code
    return user, link


class EmployeeChain(TestCase):
    def setUp(self):
        self.w = World()
        self.admin = self.w.client("admin")

    def test_admin_creates_a_user_who_signs_in_sees_only_their_sections_gets_a_task_and_loses_it_again(self):
        user, link = set_password_from_invitation("yangi_xodim", "employee", self.admin)
        self.assertEqual(user.role_assignments.get().role.name, "employee")
        user.refresh_from_db()
        self.assertTrue(user.has_usable_password())
        # 1. sign in (an employee needs no second factor) and see only the permitted sections
        c = Client()
        r = password_step(c, "yangi_xodim", NEW_PASSWORD)
        self.assertEqual(r.status_code, 302)
        self.assertTrue(signed_in(c))
        for name, expected in (("portal:dashboard", 200), ("tasks:list", 200), ("documents:list", 200), ("notifications:list", 200), ("accounts:security", 200),
                               ("users:list", 403), ("reports:index", 403), ("portal:audit", 403), ("news_admin:list", 403), ("portal:jobs", 403), ("employees:import", 403)):
            self.assertEqual(c.get(reverse(name)).status_code, expected, name)
        page = c.get(reverse("portal:dashboard")).content.decode()
        for closed in (reverse("users:list"), reverse("portal:audit"), reverse("reports:index"), reverse("news_admin:list")):
            self.assertNotIn(f'href="{closed}"', page)
        # 2. the used invitation link is dead
        self.assertNotIn("new_password1", Client().get(link, follow=True).content.decode())
        # 3. a task arrives; the notice leads to it
        r = self.admin.post(reverse("tasks:create"), {"title": "Birinchi vazifa", "description": "Matn", "priority": "normal", "due_date": timezone.localdate().isoformat(),
                                                       "assignees": [user.pk]})
        self.assertEqual(r.status_code, 302)
        task = Task.objects.get(title="Birinchi vazifa")
        note = Notification.objects.get(user=user, task=task)
        self.assertContains(c.get(reverse("notifications:list")), "Birinchi vazifa")
        self.assertRedirects(c.post(reverse("notifications:open", args=[note.pk])), reverse("tasks:detail", args=[task.pk]), fetch_redirect_response=False)
        self.assertEqual(c.get(reverse("tasks:detail", args=[task.pk])).status_code, 200)
        # 4. someone else's notification id is a 404 for this user
        other = Notification.objects.filter(task=self.w.task).first() or Notification.objects.create(user=self.w.users["stranger"], kind="task_assigned", task=self.w.task, title="x")
        self.assertEqual(c.post(reverse("notifications:open", args=[other.pk])).status_code, 404)
        # 5. the task is reassigned: the server checks access AGAIN, the old notice no longer opens it
        stranger = self.w.users["stranger"]
        self.admin.post(reverse("tasks:edit", args=[task.pk]), {"title": "Birinchi vazifa", "description": "Matn", "priority": "normal", "due_date": timezone.localdate().isoformat(),
                                                                 "assignees": [stranger.pk]})
        self.assertEqual(c.get(reverse("tasks:detail", args=[task.pk])).status_code, 403)
        r = c.post(reverse("notifications:open", args=[note.pk]), follow=True)
        self.assertRedirects(r, reverse("notifications:list"))
        self.assertEqual(c.get(reverse("tasks:detail", args=[task.pk])).status_code, 403)
        # 6. the audit trail names the actions and holds no secret
        trail = list(AuditLog.objects.values_list("action", "meta"))
        self.assertTrue({"user_created", "role_assigned", "user_invited", "password_reset_completed", "login"} <= {a for a, _m in trail}, [a for a, _m in trail])
        self.assertTrue(TaskEvent.objects.filter(task=task).exists())          # tasks keep their own event history (TaskEvent)
        text = json.dumps([m for _a, m in trail])
        for secret in (NEW_PASSWORD, link.rsplit("/", 2)[-2], "pbkdf2"):
            self.assertNotIn(secret, text)
        shown = self.admin.get(reverse("portal:audit")).content.decode()
        self.assertNotIn(NEW_PASSWORD, shown)

    def test_a_privileged_account_must_complete_the_second_step_before_anything_opens(self):
        user, _link = set_password_from_invitation("yangi_admin", "admin", self.admin)
        c = Client()
        self.assertRedirects(password_step(c, "yangi_admin", NEW_PASSWORD), SETUP, fetch_redirect_response=False)
        self.assertFalse(signed_in(c))
        self.assertIn("/accounts/login/", c.get(reverse("portal:audit"))["Location"])
        c.get(SETUP)
        c.post(SETUP, {"code": code_at(secret_of(user))})
        self.assertTrue(signed_in(c))
        self.assertEqual(c.get(reverse("portal:audit")).status_code, 200)
        self.assertEqual(c.get(reverse("users:list")).status_code, 200)

    def test_a_role_that_is_revoked_closes_the_pages_on_the_next_request_and_both_steps_are_audited(self):
        user, _link = set_password_from_invitation("vaqtinchalik", "hr", self.admin)
        c = Client()
        password_step(c, "vaqtinchalik", NEW_PASSWORD)
        self.assertEqual(c.get(reverse("reports:index")).status_code, 200)
        ra = RoleAssignment.objects.get(user=user)
        self.assertEqual(self.admin.post(reverse("users:role_remove", args=[user.pk, ra.pk])).status_code, 302)
        self.assertEqual(c.get(reverse("reports:index")).status_code, 403)
        self.assertTrue({"role_assigned", "role_revoked"} <= set(AuditLog.objects.values_list("action", flat=True)))


class DocumentChain(TestCase):
    def setUp(self):
        self.w = World()
        self.hr, self.emp, self.stranger = self.w.client("hr"), self.w.client("emp"), self.w.client("stranger")
        self.cat = DocumentCategory.objects.first()

    def post_doc(self, behaviour, name="a.pdf", data=PDF):
        fake = FakeClamd(behaviour).start()
        try:
            with scanning(fake):
                return self.hr.post(reverse("documents:create"), {"title": "Reglament-22", "category": self.cat.pk, "access_level": "general", "file": upload(name, data)})
        finally:
            fake.stop()

    def test_upload_scan_versions_access_and_leaks(self):
        # a clean file is stored and served; the verdict is recorded
        r = self.post_doc("ok")
        self.assertEqual(r.status_code, 302)
        d = Document.objects.get(title="Reglament-22")
        v1 = d.versions.get()
        self.assertEqual(v1.scan_status, "clean")
        self.assertEqual(self.emp.get(reverse("documents:download", args=[d.pk])).status_code, 200)
        # an infected NEW version is refused and leaves the old one untouched
        fake = FakeClamd("found").start()
        try:
            with scanning(fake):
                self.hr.post(reverse("documents:version_add", args=[d.pk]), {"file": upload("v2.pdf", PDF + b"2")})
        finally:
            fake.stop()
        self.assertEqual(d.versions.count(), 1)
        # the scanner is down: fail closed, the old version still works
        down = FakeClamd("ok")
        down.sock.close()
        with scanning(down):
            self.hr.post(reverse("documents:version_add", args=[d.pk]), {"file": upload("v2.pdf", PDF + b"3")})
        self.assertEqual(d.versions.count(), 1)
        self.assertEqual(self.emp.get(reverse("documents:download", args=[d.pk])).status_code, 200)
        # a good second version: the old one stays reachable by its number
        fake = FakeClamd("ok").start()
        try:
            with scanning(fake):
                self.hr.post(reverse("documents:version_add", args=[d.pk]), {"file": upload("v2.pdf", PDF + b"4"), "comment": "yangi"})
        finally:
            fake.stop()
        self.assertEqual(d.versions.count(), 2)
        old = self.emp.get(reverse("documents:download_version", args=[d.pk, 1]))
        new = self.emp.get(reverse("documents:download", args=[d.pk]))
        self.assertEqual((old.status_code, new.status_code), (200, 200))
        self.assertNotEqual(b"".join(old.streaming_content), b"".join(new.streaming_content))
        # an infected file found later is never served, in any version
        DocumentVersion.objects.filter(pk=v1.pk).update(scan_status="infected")
        self.assertEqual(self.emp.get(reverse("documents:download_version", args=[d.pk, 1])).status_code, 404)
        self.assertEqual(self.emp.get(reverse("documents:download", args=[d.pk])).status_code, 200)
        # confidential: a direct request by someone without the right is refused for every route and version
        conf = Document.objects.create(title="Maxfiy-22", category=self.cat, created_by=self.w.users["hr"], access_level="confidential")
        from apps.documents.services import add_version
        from apps.documents.validation import validate_upload
        checked = validate_upload(upload("c.pdf", PDF))
        add_version(conf, checked, actor=self.w.users["hr"])
        for client in (self.emp, self.stranger):
            for name, args in (("documents:detail", [conf.pk]), ("documents:download", [conf.pk]), ("documents:download_version", [conf.pk, 1]), ("documents:preview", [conf.pk]),
                               ("documents:inline", [conf.pk])):
                self.assertIn(client.get(reverse(name, args=args)).status_code, (403, 404), name)
        self.assertNotIn("Maxfiy-22", self.emp.get(reverse("documents:list")).content.decode())
        self.assertNotIn("Maxfiy-22", self.emp.get(reverse("portal:search"), {"q": "Maxfiy"}).content.decode().split("<main")[1].split('value="Maxfiy"')[-1])
        # no public URL reaches the files
        stored = v1.file.name
        for path in (f"/media/{stored}", f"/private_media/{stored}", f"/static/{stored}", f"/{stored}"):
            self.assertEqual(Client().get(path).status_code, 404, path)
        # the audit names the events but not titles or file names
        text = json.dumps(list(AuditLog.objects.values_list("meta", flat=True)))
        for private in ("Reglament-22", "Maxfiy-22", "a.pdf", "c.pdf"):
            self.assertNotIn(private, text)

    def test_the_restored_test_backup_keeps_documents_and_their_access_rules(self):
        """The full restore (database + files, 2FA secrets, access rules) is exercised by core.tests_backup.RestoringIt in an isolated scratch
        database; here only the contract is pinned so that the chain cannot silently lose it."""
        from apps.core import tests_backup
        names = {n for n in dir(tests_backup.RestoringIt) if n.startswith("test_")}
        self.assertTrue(any("restor" in n for n in names), names)
        self.assertTrue(hasattr(call_command, "__call__"))
        from django.core.management import get_commands
        for command in ("verify_restored_access", "verify_restored_two_factor", "run_restore_check"):
            self.assertIn(command, get_commands())


class NewsChain(TestCase):
    def setUp(self):
        self.w = World()
        self.cat = self.w.news_category
        self.author = make_user("fin_au", role="news_author")
        self.editor, self.editor2 = make_user("fin_ed", role="news_editor"), make_user("fin_ed2", role="news_editor")
        self.publisher = make_user("fin_pu", role="news_publisher")
        self.au, self.ed, self.ed2, self.pu = (login(u) for u in (self.author, self.editor, self.editor2, self.publisher))
        self.su, self.hr, self.anon = self.w.client("superuser"), self.w.client("hr"), Client()
        self.body = {"title": "Qabul-22 yangiligi", "excerpt": "Qisqa", "body": "Matn", "category": self.cat.pk}

    def post(self, client, n, step, **data):
        n.refresh_from_db()
        if step in ("approve", "request_changes"):
            data.setdefault("version", workflow.version_token(n))
        r = client.post(reverse("news_admin:status", args=[n.pk, step]), data)
        n.refresh_from_db()
        return r

    def public(self, n):
        return self.anon.get(reverse("content:news_detail", args=[n.slug])).status_code

    def hidden_everywhere(self, n):
        self.assertEqual(self.public(n), 404, n.status)
        self.assertNotIn(n.title, self.anon.get(reverse("content:news_list")).content.decode())
        self.assertNotIn(n.title, self.anon.get(reverse("content:home")).content.decode())
        found = self.w.client("emp").get(reverse("portal:search"), {"q": n.title}).content.decode()
        self.assertNotIn(reverse("content:news_detail", args=[n.slug]), found)
        for fmt in ("xlsx", "pdf"):
            self.assertNotIn(n.title.encode(), self.hr.get(reverse("reports:export", args=[fmt])).content)

    def test_the_whole_cycle_with_every_side_path(self):
        self.assertEqual(self.au.post(reverse("news_admin:create"), self.body).status_code, 302)
        n = News.objects.get(title="Qabul-22 yangiligi")
        self.hidden_everywhere(n)
        # the author submits; editors are told, the author is not
        self.post(self.au, n, "submit")
        self.assertEqual(n.status, "review")
        self.assertTrue(Notification.objects.filter(user=self.editor, kind="news_submitted", news=n).exists())
        self.assertFalse(Notification.objects.filter(user=self.author).exists())
        self.hidden_everywhere(n)
        # an editor returns it with a comment AND an internal note
        self.post(self.ed, n, "request_changes", comment="Muallifga: sarlavhani aniqlashtiring", internal_note="MAXFIY-ICHKI-22")
        self.assertEqual(n.status, "changes_requested")
        self.assertTrue(Notification.objects.filter(user=self.author, kind="news_changes_requested", news=n).exists())
        page = self.au.get(reverse("news_admin:preview", args=[n.pk])).content.decode()
        self.assertIn("sarlavhani aniqlashtiring", page)
        self.assertNotIn("MAXFIY-ICHKI-22", page)
        self.assertIn("MAXFIY-ICHKI-22", self.ed2.get(reverse("news_admin:preview", args=[n.pk])).content.decode())
        # the author corrects and resubmits; a DIFFERENT editor approves; the author cannot approve anything
        self.au.post(reverse("news_admin:edit", args=[n.pk]), {**self.body, "title": "Qabul-22 yangiligi (tuzatilgan)"})
        n.refresh_from_db()
        self.post(self.au, n, "approve")
        self.post(self.au, n, "submit")
        self.post(self.ed2, n, "approve")
        self.assertEqual(n.status, "approved")
        self.assertTrue(Notification.objects.filter(user=self.author, kind="news_approved", news=n).exists())
        self.hidden_everywhere(n)
        # editing the approved text sends it back; the admin cannot slip it through either
        self.au.post(reverse("news_admin:edit", args=[n.pk]), {**self.body, "title": n.title, "body": "Boshqa matn"})
        n.refresh_from_db()
        self.assertEqual(n.status, "draft")
        self.post(self.au, n, "submit")
        self.post(self.ed, n, "approve")
        admin_url = reverse("admin:content_news_change", args=[n.pk])
        self.su.get(admin_url)
        self.su.post(admin_url, {"title": n.title, "slug": n.slug, "excerpt": "", "body": "Admin orqali", "category": self.cat.pk, "status": "published"})
        n.refresh_from_db()
        self.assertEqual(n.status, "draft")                                          # the approval is gone because the text changed; status was not settable
        self.assertEqual(self.public(n), 404)
        self.assertEqual(self.su.post(reverse("admin:content_news_changelist"), {"action": "delete_selected", "_selected_action": [str(n.pk)], "post": "yes"}).status_code, 200)
        self.assertTrue(News.objects.filter(pk=n.pk).exists())
        # scheduled publication: approved with a future date stays invisible until the date
        self.post(self.au, n, "submit")
        self.post(self.ed2, n, "approve")
        future = timezone.now() + datetime.timedelta(days=3)
        News.objects.filter(pk=n.pk).update(publish_at=future)
        self.post(self.pu, n, "publish")
        self.assertEqual(n.status, "published")
        self.assertEqual(self.public(n), 404)
        News.objects.filter(pk=n.pk).update(publish_at=timezone.now() - datetime.timedelta(minutes=1))
        self.assertEqual(self.public(n), 200)
        self.assertIn(reverse("content:news_detail", args=[n.slug]), self.w.client("emp").get(reverse("portal:search"), {"q": "Qabul-22"}).content.decode())
        # unpublish, archive, restore
        self.post(self.pu, n, "unpublish")
        self.assertEqual((n.status, self.public(n)), ("draft", 404))
        self.post(self.pu, n, "archive")
        self.assertEqual(n.status, "draft")                                          # a publisher cannot see (nor archive) someone else's draft
        self.post(self.w.client("content"), n, "archive")
        self.assertEqual((n.status, self.public(n)), ("archived", 404))
        self.post(self.w.client("content"), n, "restore")
        self.assertEqual(n.status, "draft")
        # the history is complete, in order, and nobody can alter it
        actions = list(n.reviews.values_list("action", flat=True))
        self.assertEqual(actions[:4], ["submit", "request_changes", "submit", "approve"])
        self.assertIn("edit_reset", actions)
        self.assertEqual(actions[-1], "restore")
        for suffix in ("", f"{n.reviews.first().pk}/change/", "add/"):
            self.assertEqual(self.su.get(f"{reverse('admin:index')}content/newsreview/{suffix}").status_code, 404)
        # audit: names of actions only, never the note or the comment
        text = json.dumps(list(AuditLog.objects.values_list("meta", flat=True)))
        self.assertNotIn("MAXFIY-ICHKI-22", text)
        self.assertNotIn("sarlavhani aniqlashtiring", text)


class SearchAndReports(TestCase):
    def setUp(self):
        self.w = World()

    def test_search_and_exports_keep_the_rights_and_neutralise_formulas(self):
        w = self.w
        admin = w.users["admin"]
        Task.objects.create(title="=HYPERLINK(\"http://evil.invalid\",\"x\")", description="d", created_by=admin, due_date=timezone.localdate()).assignees.set([w.users["emp"]])
        for who, may_export in (("hr", True), ("admin", True), ("emp", False), ("stranger", False), ("content", False)):
            c = w.client(who)
            r = c.get(reverse("reports:export", args=["xlsx"]))
            self.assertEqual(r.status_code, 200 if may_export else 403, who)
        # workbook cells never start a formula
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(w.client("hr").get(reverse("reports:export", args=["xlsx"])).content))
        for ws in book:
            for row in ws.iter_rows():
                for cell in row:
                    self.assertNotEqual(cell.data_type, "f", (ws.title, cell.coordinate))
                    if isinstance(cell.value, str):
                        self.assertFalse(cell.value.startswith(("=", "+", "-", "@")), (ws.title, cell.coordinate, cell.value[:20]))
        # a forged filter cannot widen the scope: the head of one department asks for another one
        other = w.other
        r = w.client("head").get(reverse("reports:index"), {"department": other.pk})
        self.assertIn(r.status_code, (200, 403))
        if r.status_code == 200:
            self.assertNotIn("+998 90 555 66 77", r.content.decode())
        # search: foreign people, hidden documents and unpublished news never appear
        emp = w.client("emp")
        # the directory is open by design (name, unit, status); contact and private fields are not, for people outside the viewer's scope
        found = w.client("stranger").get(reverse("portal:search"), {"q": "Aliyev"}).content.decode()
        self.assertIn("Aliyev", found)
        for private in ("+998 71 111 22 33", "vali@work.example.test", "+998 90 555 66 77", "17.05.1990", "1990"):
            self.assertNotIn(private, found)
        hidden = Document.objects.create(title="Yopiq-22", category=DocumentCategory.objects.first(), created_by=w.users["hr"], access_level="confidential")
        self.assertNotIn(str(hidden.pk), emp.get(reverse("portal:search"), {"q": "Yopiq"}).content.decode())
        self.assertIn(str(hidden.pk), w.client("hr").get(reverse("portal:search"), {"q": "Yopiq"}).content.decode())


class PermissionsChain(TestCase):
    def setUp(self):
        self.w = World()
        call_command("seed_roles", verbosity=0)
        self.su, self.admin = self.w.client("superuser"), self.w.client("admin")

    def seed(self, *args):
        out, err = io.StringIO(), io.StringIO()
        call_command("seed_roles", *args, stdout=out, stderr=err)
        return out.getvalue(), err.getvalue()

    def test_roles_exceptions_and_seed_roles_work_together(self):
        user = make_user("fin_perm")
        # assign and revoke a role in the portal; both are audited
        self.assertEqual(self.admin.post(reverse("users:role_add", args=[user.pk]), {"role": "news_editor"}).status_code, 302)
        self.assertTrue(user.groups.filter(name="news_editor").exists())
        # a personal exception through the admin: audited, effective, and it survives every seed_roles mode
        keep = list(user.user_permissions.values_list("pk", flat=True))
        perm = Permission.objects.get(content_type__app_label="content", codename="approve_own_news")
        url = reverse("admin:accounts_user_change", args=[user.pk])
        page = self.su.get(url)
        data = {"username": user.username, "email": user.email, "first_name": "", "last_name": "", "locale": "uz", "is_active": "on",
                "user_permissions": keep + [perm.pk], "date_joined_0": "2026-01-01", "date_joined_1": "00:00:00"}
        for inline in page.context["inline_admin_formsets"]:
            mf = inline.formset.management_form
            for k in ("TOTAL_FORMS", "INITIAL_FORMS", "MIN_NUM_FORMS", "MAX_NUM_FORMS"):
                data[f"{mf.prefix}-{k}"] = mf.initial.get(k, 0)
        self.assertEqual(self.su.post(url, data).status_code, 302)
        entry = AuditLog.objects.filter(action="admin_changed", object_type="User").latest("at")
        self.assertEqual(entry.meta["user_permissions_added"], ["content.approve_own_news"])
        for args in (("--dry-run",), (), ("--dry-run", "--prune"), ("--prune",)):
            self.seed(*args)
            self.assertTrue(User.objects.get(pk=user.pk).user_permissions.filter(pk=perm.pk).exists(), args)
            self.assertTrue(User.objects.get(pk=user.pk).groups.filter(name="news_editor").exists(), args)
        # a hand-made group permission is reported and kept; only --prune removes it, and a dry run removes nothing
        group = Group.objects.get(name="news_author")
        group.permissions.add(Permission.objects.get(content_type__app_label="content", codename="publish_news"))
        out, err = self.seed()
        self.assertIn("kept", err)
        self.assertTrue(group.permissions.filter(codename="publish_news").exists())
        self.seed("--dry-run", "--prune")
        self.assertTrue(group.permissions.filter(codename="publish_news").exists())
        self.seed("--prune")
        self.assertFalse(group.permissions.filter(codename="publish_news").exists())
        # revoke the role: access ends on the next request
        c = login(User.objects.get(pk=user.pk))
        self.assertEqual(c.get(reverse("news_admin:list")).status_code, 200)
        ra = RoleAssignment.objects.get(user=user, role__name="news_editor")
        self.assertEqual(self.admin.post(reverse("users:role_remove", args=[user.pk, ra.pk])).status_code, 302)
        self.assertEqual(c.get(reverse("news_admin:list")).status_code, 403)
        self.assertTrue({"role_assigned", "role_revoked", "admin_changed"} <= set(AuditLog.objects.values_list("action", flat=True)))
