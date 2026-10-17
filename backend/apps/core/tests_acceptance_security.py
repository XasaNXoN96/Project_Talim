"""Security acceptance: XSS through every user-entered field, open redirects, junk input, direct file URLs, and personal-data
leaks through audit, notifications, e-mails, reports and exports. Fictional data only."""
import io
import json
import random
import re
from pathlib import Path

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core import mail
from django.test import Client, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounts.testing import PASSWORD
from apps.audit.models import AuditLog
from apps.documents.models import Document, DocumentCategory
from apps.documents.testing import PDF, upload
from apps.notifications.models import Notification
from apps.organization.models import Department, Employee, Position
from apps.tasks.models import Task, TaskEvent

from .acceptance import ACTORS, World, app_routes

User = get_user_model()
PAYLOADS = ('<script>alert("XSS1")</script>', '"><img src=x onerror=alert(2)>', "'><svg/onload=alert(3)>", "</textarea><script>alert(4)</script>")
RAW_MARKERS = ('<script>alert(', "<img src=x onerror", "<svg/onload", "</textarea><script")


class XssTests(TestCase):
    def test_user_entered_text_is_escaped_on_every_page_that_shows_it(self):
        w = World()
        p = PAYLOADS[0]
        q = PAYLOADS[1]
        r = PAYLOADS[2]
        emp = w.users["emp"]
        emp.first_name, emp.last_name = p, q
        emp.save()
        Department.objects.filter(pk=w.child.pk).update(name=p)
        Position.objects.filter(pk=w.position.pk).update(name=q)
        Employee.objects.filter(pk=w.card.pk).update(last_name=r, first_name=p, work_phone=q, work_email="x@example.test")
        Task.objects.filter(pk=w.task.pk).update(title=p, description=q)
        w.task.comments.create(author=emp, body=r)
        Document.objects.filter(pk=w.doc.pk).update(title=p, description=q, number=r)
        cat = DocumentCategory.objects.create(name=q)
        Document.objects.filter(pk=w.doc.pk).update(category=cat)
        from apps.tasks.models import TaskAttachment

        TaskAttachment.objects.filter(pk=w.attachment.pk).update(original_name=p + ".pdf")
        w.doc.versions.update(original_name=q + ".pdf", comment=r)
        Notification.objects.filter(pk=w.notification.pk).update(title=p)
        w.news.title, w.news.body, w.news.excerpt = p, q + "\n" + r, r
        w.news.save()
        pages = [
            ("users:list", {}), ("users:detail", {"pk": emp.pk}), ("users:edit", {"pk": emp.pk}), ("employees:list", {}), ("employees:detail", {"pk": w.card.pk}),
            ("employees:edit", {"pk": w.card.pk}), ("employees:structure", {}), ("employees:department_edit", {"pk": w.child.pk}), ("employees:position_edit", {"pk": w.position.pk}),
            ("tasks:list", {}), ("tasks:detail", {"pk": w.task.pk}), ("tasks:edit", {"pk": w.task.pk}), ("tasks:create", {}),
            ("documents:list", {}), ("documents:detail", {"pk": w.doc.pk}), ("documents:edit", {"pk": w.doc.pk}), ("documents:create", {}),
            ("notifications:list", {}), ("reports:index", {}), ("portal:dashboard", {}), ("portal:audit", {}),
            ("content:home", {}), ("content:news_list", {}), ("content:news_detail", {"slug": w.news.slug}),
        ]
        c = w.client("superuser")
        shown = 0
        for name, kw in pages:
            r_ = c.get(reverse(name, kwargs=kw), {"date_from": "2020-01-01", "tab": "all"})
            self.assertEqual(r_.status_code, 200, name)
            body = r_.content.decode()
            for m in RAW_MARKERS:
                self.assertNotIn(m, body, f"{name}: raw markup reached the page ({m})")
            if "&lt;script&gt;" in body or "&lt;img src=x" in body or "&lt;svg/onload" in body:
                shown += 1
        self.assertGreaterEqual(shown, 15, "the payloads were not displayed anywhere: the test would prove nothing")

    def test_payloads_in_exports_and_e_mails_are_inert(self):
        w = World()
        Department.objects.filter(pk=w.child.pk).update(name='=HYPERLINK("http://evil.test","x")')
        Task.objects.filter(pk=w.task.pk).update(title='<script>alert(1)</script>')
        wb = load_workbook(io.BytesIO(w.client("admin").get(reverse("reports:export", args=["xlsx"])).content))
        for ws in wb:
            for row in ws.iter_rows():
                for cell in row:
                    self.assertNotEqual(cell.data_type, "f")
                    if isinstance(cell.value, str):
                        self.assertFalse(cell.value.startswith(("=", "+", "-", "@")), cell.value)
        emp = w.users["emp"]
        emp.first_name = "<b>Vali</b>"
        emp.save()
        w.client("admin").post(reverse("users:action", args=[emp.pk, "invite"]))
        self.assertEqual(mail.outbox[-1].content_subtype, "plain")  # plain-text mail: nothing is interpreted as HTML

    def test_json_and_header_injection_in_filenames(self):
        w = World()
        c = w.client("hr")
        cat = DocumentCategory.objects.first()
        evil = 'x"\r\nX-Injected: yes.pdf'
        r = c.post(reverse("documents:create"), {"title": 'T"\r\nX-Injected: 1', "category": cat.pk, "access_level": "general", "file": upload(evil, PDF)})
        self.assertLess(r.status_code, 500)
        doc = Document.objects.filter(title__startswith="T").first()
        if doc:
            d = c.get(reverse("documents:download", args=[doc.pk]))
            self.assertNotIn("X-Injected", d.headers)
            self.assertNotIn("\n", d["Content-Disposition"])


class OpenRedirectTests(TestCase):
    def setUp(self):
        self.w = World()
        self.evil = ("https://evil.test/x", "//evil.test/x", "/\\evil.test", "javascript:alert(1)", "http:evil.test", "///evil.test", "https:///evil.test")

    def test_login_next_is_validated_for_plain_and_for_two_factor_users(self):
        for evil in self.evil:
            c = Client()
            r = c.post(reverse("accounts:login") + "?next=" + evil, {"username": "emp", "password": PASSWORD, "next": evil})
            self.assertEqual(r.status_code, 302)
            self.assertFalse(r["Location"].startswith(("http", "//", "javascript", "/\\", "///")), (evil, r["Location"]))
        from apps.accounts.tests_twofactor import enrol

        admin = self.w.users["admin"]
        secret, _codes = enrol(admin)
        from apps.accounts.tests_twofactor import code_at

        for i, evil in enumerate(self.evil[:3]):
            c = Client()
            r = c.post(reverse("accounts:login"), {"username": "boss", "password": PASSWORD, "next": evil})
            self.assertEqual(r.status_code, 302)
            r = c.post(reverse("accounts:two_factor_verify"), {"code": code_at(secret, i - 1)})
            self.assertEqual(r.status_code, 302)
            self.assertFalse(r["Location"].startswith(("http", "//", "javascript")), (evil, r["Location"]))

    def test_good_internal_next_still_works(self):
        c = Client()
        r = c.post(reverse("accounts:login"), {"username": "emp", "password": PASSWORD, "next": "/portal/tasks/?tab=all"})
        self.assertEqual(r["Location"], "/portal/tasks/?tab=all")

    def test_logout_and_language_switch_ignore_foreign_next(self):
        c = self.w.client("emp")
        for evil in self.evil:
            r = c.post(reverse("set_language"), {"language": "ru", "next": evil})
            self.assertFalse(r["Location"].startswith(("http", "//", "javascript", "/\\", "///")), (evil, r["Location"]))
        c2 = self.w.client("emp")
        r = c2.post(reverse("accounts:logout") + "?next=https://evil.test/")
        self.assertFalse(r["Location"].startswith("http"), r["Location"])


class RobustnessTests(TestCase):
    """Junk in, no 5xx out: every POST route, several hostile payloads, an authorised actor."""

    def test_hostile_bodies_never_crash_any_post_route(self):
        w = World()
        random.Random(7)
        payloads = [
            {}, {"title": "x" * 20000, "body": "y" * 50000}, {"pk": "not-a-uuid", "assignees": ["zzz", "00000000-0000-0000-0000-000000000000"]},
            {"date_from": "2020-13-45", "date_to": "9999-99-99", "department": "'; DROP TABLE accounts_user; --", "category": "../../etc/passwd"},
            {"username": "a\x00b", "email": "not-an-email", "role": "'\"<>", "department": "x"}, {"file": "plain text, not a file", "language": "ru\x00"},
            {"assignees": [str(w.users["emp"].pk)] * 500}, {f"k{i}": "v" * 50 for i in range(300)},
        ]
        c = w.client("superuser")
        failures = []
        for name, _cls in app_routes():
            if name in ("accounts:logout", "healthz"):
                continue
            kw = w.kwargs_for(name)
            for i, body in enumerate(payloads):
                r = c.post(reverse(name, kwargs=kw), body)
                if r.status_code >= 500:
                    failures.append(f"{name} payload#{i} -> {r.status_code}")
        self.assertEqual(failures, [])

    def test_hostile_uploads_never_crash_upload_routes(self):
        w = World()
        c = w.client("admin")
        cat = DocumentCategory.objects.first()
        junk = [("", b""), ("a.pdf", b""), ("a.pdf", bytes(rnd) if (rnd := random.Random(3).randbytes(200_000)) else b""), ("../../../etc/passwd", PDF), ("a" * 400 + ".pdf", PDF),
                ("a\x00.pdf", PDF), ("a.pdf\r\nX: y", PDF), ("‮gnp.pdf", PDF), ("CON.pdf", PDF), (".pdf", PDF), ("a.PDF ", PDF)]
        for fname, data in junk:
            for name, kw, extra in (
                ("documents:create", {}, {"title": "T", "category": cat.pk, "access_level": "general"}),
                ("documents:version_add", {"pk": w.doc.pk}, {"comment": "c"}),
                ("tasks:attachment_add", {"pk": w.task.pk}, {}),
            ):
                try:
                    f = upload(fname, data)
                except Exception:  # noqa: BLE001
                    continue
                r = c.post(reverse(name, kwargs=kw), {**extra, "file": f})
                self.assertLess(r.status_code, 500, (name, fname[:20]))
        # nothing was written under a client-chosen path
        root = Path(settings.PRIVATE_MEDIA_ROOT).resolve()
        for p in root.rglob("*"):
            if p.is_file():
                self.assertRegex(str(p.relative_to(root)), r"^[0-9a-f]{2}/[0-9a-f]{32}\.[a-z0-9]+$")

    def test_error_pages_do_not_reveal_internals(self):
        w = World()
        c = Client(raise_request_exception=False)
        r = c.get("/no-such/")
        self.assertEqual(r.status_code, 404)
        for leak in ("Traceback", "django.", "SECRET_KEY", "settings.py", "/home/", "psycopg"):
            self.assertNotIn(leak, r.content.decode())
        from unittest import mock

        c2 = w.client("admin")
        c2.raise_request_exception = False
        with mock.patch("apps.reports.views.build_report", side_effect=RuntimeError("password=hunter2 /home/secret/path")):
            r = c2.get(reverse("reports:index"))
        self.assertEqual(r.status_code, 500)
        for leak in ("hunter2", "/home/secret", "Traceback", "RuntimeError"):
            self.assertNotIn(leak, r.content.decode())


class DirectFileAccessTests(TestCase):
    def test_stored_files_have_no_url_reachable_by_anyone(self):
        w = World()
        names = [w.attachment.file.name, w.doc.versions.get().file.name, w.secret_doc.versions.get().file.name]
        prefixes = ["/media/", "/static/", "/private_media/", "/portal/documents/", "/portal/tasks/", "/", "/portal/", "/documents/", "/files/", "/uploads/", "/media/../", "/static/../"]
        for actor in ACTORS:
            c = w.client(actor)
            for name in names:
                for pre in prefixes:
                    for variant in (name, "../" + name, name.replace("/", "%2F"), name.upper()):
                        r = c.get(pre + variant)
                        self.assertIn(r.status_code, (404, 302, 403, 400), (actor, pre + variant, r.status_code))
                        self.assertNotEqual(r["Content-Type"] if r.has_header("Content-Type") else "", "application/pdf", (actor, pre + variant))

    def test_file_responses_cannot_be_rendered_by_the_browser(self):
        w = World()
        for url in (reverse("documents:download", args=[w.doc.pk]), reverse("tasks:attachment_download", args=[w.task.pk, w.attachment.pk])):
            r = w.client("admin").get(url)
            self.assertEqual(r["Content-Type"], "application/octet-stream")
            self.assertTrue(r["Content-Disposition"].startswith("attachment"))
            self.assertEqual(r["X-Content-Type-Options"], "nosniff")
            self.assertIn("sandbox", r["Content-Security-Policy"])
            self.assertIn("no-store", r["Cache-Control"])


class PersonalDataLeakTests(TestCase):
    SECRETS = {
        "personal_phone": "+998 90 555 66 77", "birth": "17.05.1990", "birth_iso": "1990-05-17", "work_phone": "+998 71 111 22 33",
        "work_email": "vali@work.example.test", "emp_email": "emp@example.test", "stranger_email": "stranger@example.test",
        "first": "Vali", "last": "Aliyev", "password": PASSWORD, "fname": "Maxfiy-shaxsiy-fayl",
    }

    def setUp(self):
        self.w = World()

    def drive_the_system(self):
        """Exercise most state-changing paths once so every kind of audit entry exists."""
        w = self.w
        admin, hr = w.client("admin"), w.client("hr")
        anon = Client()
        anon.post(reverse("accounts:login"), {"username": "emp@example.test", "password": "wrong"})  # an e-mail typed as identifier
        anon.post(reverse("accounts:login"), {"username": "no-such-user", "password": "wrong"})
        anon.post(reverse("accounts:password_reset"), {"email": "emp@example.test"})
        anon.post(reverse("accounts:password_reset"), {"email": "unknown@example.test"})
        admin.post(reverse("users:create"), {"username": "newbie", "email": "newbie.private@example.test", "last_name": "Yangiyev", "first_name": "Anvar", "role": "employee"})
        newbie = User.objects.get(username="newbie")
        admin.post(reverse("users:edit", args=[newbie.pk]), {"email": "changed.private@example.test", "last_name": "Boshqayev", "first_name": "Anvar", "locale": "ru"})
        admin.post(reverse("users:action", args=[newbie.pk, "invite"]))
        admin.post(reverse("users:action", args=[newbie.pk, "force_password"]))
        admin.post(reverse("users:action", args=[newbie.pk, "deactivate"]))
        hr.post(reverse("employees:edit", args=[w.card.pk]), {
            "last_name": "Aliyev", "first_name": "Vali", "department": w.child.pk, "position": w.position.pk, "status": "active", "work_phone": "+998 71 999 00 11",
            "work_email": "vali.new@work.example.test", "personal_phone": "+998 90 123 45 67", "birth_date": "1991-06-18", "hired_at": "2020-01-01"})
        hr.get(reverse("employees:detail", args=[w.card.pk]))
        hr.post(reverse("documents:create"), {"title": "Sarlavha", "category": DocumentCategory.objects.first().pk, "access_level": "confidential",
                                              "file": upload("Maxfiy-shaxsiy-fayl.pdf", PDF)})
        d = Document.objects.get(title="Sarlavha")
        hr.get(reverse("documents:detail", args=[d.pk]))
        hr.get(reverse("documents:download", args=[d.pk]))
        hr.post(reverse("documents:archive", args=[d.pk]))
        emp = w.client("emp")
        emp.post(reverse("tasks:attachment_add", args=[w.task.pk]), {"file": upload("Maxfiy-shaxsiy-fayl.pdf", PDF + b"1")})
        emp.post(reverse("tasks:comment", args=[w.task.pk]), {"body": "Shaxsiy izoh: Vali Aliyev +998 90 555 66 77"})
        emp.post(reverse("tasks:transition", args=[w.task.pk, "done"]))
        admin.post(reverse("tasks:transition", args=[w.task.pk, "done"]))
        admin.get(reverse("reports:index"))
        admin.get(reverse("reports:export", args=["xlsx"]))
        admin.get(reverse("reports:export", args=["pdf"]))

    def test_audit_log_never_contains_personal_data_secrets_or_content(self):
        self.drive_the_system()
        forbidden = ["+998", "1990", "1991", "@example.test", "@work.example", "Aliyev", "Yangiyev", "Boshqayev", "Maxfiy-shaxsiy", "Shaxsiy izoh", PASSWORD, "hunter2"]
        rows = list(AuditLog.objects.all())
        self.assertGreater(len(rows), 20)
        for e in rows:
            blob = json.dumps(e.meta, ensure_ascii=False) + e.object_id + e.object_type + e.get_action_display()
            for f in forbidden:
                self.assertNotIn(f, blob, f"{e.action}: '{f}' in {blob[:200]}")
        # failed logins keep a hash, never the typed identifier
        failed = AuditLog.objects.filter(action=AuditLog.Action.LOGIN_FAILED)
        self.assertTrue(failed.exists())
        for e in failed:
            self.assertNotIn("emp@example.test", json.dumps(e.meta))
            self.assertNotIn("no-such-user", json.dumps(e.meta))

    def test_task_history_and_notifications_never_carry_comments_or_file_names(self):
        self.drive_the_system()
        for ev in TaskEvent.objects.all():
            blob = json.dumps(ev.meta, ensure_ascii=False)
            for f in ("Shaxsiy izoh", "Maxfiy-shaxsiy", "+998", "@example.test"):
                self.assertNotIn(f, blob)
        for n in Notification.objects.all():
            self.assertEqual(n.title, self.w.task.title)  # only the task title, nothing else
        for page in ("notifications:list",):
            body = self.w.client("emp").get(reverse(page)).content.decode()
            self.assertNotIn("Shaxsiy izoh", body)

    def test_e_mails_contain_only_what_the_recipient_needs(self):
        self.drive_the_system()
        self.assertGreater(len(mail.outbox), 1)
        for m in mail.outbox:
            blob = m.subject + m.body
            for f in ("+998", "1990", "Shaxsiy izoh", "Maxfiy-shaxsiy", PASSWORD, "birth", "personal_phone"):
                self.assertNotIn(f, blob, m.subject)
            others = [a for a in ("emp@example.test", "stranger@example.test", "newbie.private@example.test") if a not in m.to]
            for a in others:
                self.assertNotIn(a, blob, f"{m.subject}: another person's address in the body")
            self.assertEqual(len(m.to), 1)

    def test_reports_and_exports_carry_no_contact_or_private_data_for_any_role(self):
        for actor in ("head", "hr", "admin", "superuser"):
            c = self.w.client(actor)
            page = c.get(reverse("reports:index")).content.decode()
            xlsx = load_workbook(io.BytesIO(c.get(reverse("reports:export", args=["xlsx"])).content))
            cells = " ".join(str(v) for ws in xlsx for row in ws.iter_rows(values_only=True) for v in row if v is not None)
            pdf = c.get(reverse("reports:export", args=["pdf"])).content
            for label, secret in self.SECRETS.items():
                if label in ("first", "last", "password", "fname", "stranger_email", "emp_email"):
                    continue
                self.assertNotIn(secret, page, (actor, "page", label))
                self.assertNotIn(secret, cells, (actor, "xlsx", label))
                self.assertNotIn(secret.encode(), pdf, (actor, "pdf", label))
            for a in ("emp@example.test", "stranger@example.test", "@work.example"):
                self.assertNotIn(a, page + cells, (actor, a))

    def test_dashboard_notifications_and_lists_do_not_leak_other_peoples_private_fields(self):
        stranger = self.w.client("stranger")
        for name, kw in (("employees:list", {}), ("tasks:list", {}), ("documents:list", {}), ("notifications:list", {}), ("portal:dashboard", {}), ("employees:structure", {})):
            body = stranger.get(reverse(name, kwargs=kw)).content.decode()
            for label in ("personal_phone", "birth", "birth_iso", "work_phone", "work_email", "emp_email"):
                self.assertNotIn(self.SECRETS[label], body, (name, label))

    def test_password_and_token_never_appear_in_responses_or_logs(self):
        import logging

        class Capture(logging.Handler):
            def __init__(self):
                super().__init__()
                self.lines = []

            def emit(self, record):
                self.lines.append(self.format(record))

        cap = Capture()
        logging.getLogger().addHandler(cap)
        try:
            c = Client()
            c.post(reverse("accounts:password_reset"), {"email": "emp@example.test"})
            token_url = re.search(r"/accounts/reset/[^\s]+", mail.outbox[-1].body).group(0)
            c.get(token_url, follow=True)
            c.post(reverse("accounts:login"), {"username": "emp", "password": "Very-Secret-Wrong-1"})
        finally:
            logging.getLogger().removeHandler(cap)
        text = "\n".join(cap.lines)
        for secret in ("Very-Secret-Wrong-1", token_url.split("/")[-2] if token_url.count("/") > 3 else "zzzz"):
            self.assertNotIn(secret, text)
