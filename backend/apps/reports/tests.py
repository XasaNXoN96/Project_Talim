import datetime
import io
import re

from django.contrib.auth.models import Group
from django.core.management import call_command
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.access.models import RoleAssignment
from apps.accounts.testing import login_with_two_factor, make_user
from apps.audit.models import AuditLog
from apps.documents.models import Document, DocumentCategory
from apps.organization.models import Department, Employee, Position
from apps.tasks.models import Task

D = datetime.date
INDEX = reverse("reports:index")


def export(fmt):
    return reverse("reports:export", args=[fmt])


def at(day, hour=12):
    return timezone.make_aware(datetime.datetime.combine(day, datetime.time(hour)))


class Base(TestCase):
    def setUp(self):
        call_command("seed_roles", verbosity=0)
        self.today = timezone.localdate()
        self.root = Department.objects.create(name="Metodika")
        self.child = Department.objects.create(name="Maktabgacha", parent=self.root)
        self.other = Department.objects.create(name="Moliya")
        self.pos_a, self.pos_b = Position.objects.create(name="Metodist"), Position.objects.create(name="Buxgalter")

        self.admin = make_user("boss", role="admin")
        self.hr = make_user("hrista", role="hr")
        self.head = make_user("head", role="employee")
        RoleAssignment.objects.create(user=self.head, role=Group.objects.get(name="department_head"), scope="department", department=self.root)
        self.emp = make_user("plain", role="employee")

        self.a1 = self.person("a1", "Aliyev", "Vali", self.child, self.pos_a)
        self.a2 = self.person("a2", "Sodiqova", "Gulnora", self.child, self.pos_a)
        self.b1 = self.person("b1", "Botirov", "Sher", self.other, self.pos_b)
        Employee.objects.create(last_name="Arxivdagi", first_name="X", department=self.other, status="archived")
        Employee.objects.create(last_name="Dam", first_name="Y", department=self.other, position=self.pos_b, status="on_leave")

        d = datetime.timedelta
        T = Task.Status
        # inside the branch (assignees a1/a2)
        self.t_done_ontime = self.task("done on time", [self.a1], T.DONE, created=self.today - d(days=10), due=self.today - d(days=5), closed=self.today - d(days=6))
        self.t_done_late = self.task("done late", [self.a1, self.a2], T.DONE, created=self.today - d(days=9), due=self.today - d(days=5), closed=self.today - d(days=3))
        self.t_overdue = self.task("overdue", [self.a2], T.IN_PROGRESS, created=self.today - d(days=8), due=self.today - d(days=1))
        self.t_new = self.task("fresh", [self.a1], T.NEW, created=self.today - d(days=1), due=self.today + d(days=5))
        # outside the branch
        self.t_other = self.task("finance", [self.b1], T.DONE, created=self.today - d(days=4), due=self.today, closed=self.today - d(days=2))
        self.t_old = self.task("ancient", [self.a1], T.DONE, created=self.today - d(days=100), due=None, closed=self.today - d(days=90))

        self.cat1, self.cat2 = DocumentCategory.objects.get(name="Buyruqlar"), DocumentCategory.objects.get(name="Hisobotlar")
        self.doc("public1", self.cat1, "general", self.hr)
        self.doc("public2", self.cat1, "general", self.hr)
        self.doc("report1", self.cat2, "general", self.hr)
        self.doc("secret", self.cat2, "confidential", self.hr)
        self.doc("moliya-only", self.cat1, "department", self.hr, dept=self.other)

    def person(self, username, last, first, dept, pos):
        u = make_user(username, role="employee")
        u.last_name, u.first_name = last, first
        u.save()
        Employee.objects.create(last_name=last, first_name=first, department=dept, position=pos, user=u, status="active")
        return u

    def task(self, title, assignees, status, *, created, due, closed=None):
        t = Task.objects.create(title=title, created_by=self.admin, status=status, due_date=due, closed_at=at(closed) if closed else None)
        Task.objects.filter(pk=t.pk).update(created_at=at(created))
        t.assignees.set(assignees)
        return t

    def doc(self, title, cat, level, author, dept=None):
        return Document.objects.create(title=title, category=cat, created_by=author, access_level=level, department=dept)

    def client_for(self, user):
        c = Client()
        login_with_two_factor(c, user)
        return c

    def params(self, **kw):
        p = {"date_from": (self.today - datetime.timedelta(days=30)).isoformat(), "date_to": self.today.isoformat()}
        p.update({k: v for k, v in kw.items()})
        return p

    def page(self, user, **kw):
        return self.client_for(user).get(INDEX, self.params(**kw))

    def workbook(self, user, **kw):
        r = self.client_for(user).get(export("xlsx"), self.params(**kw))
        self.assertEqual(r.status_code, 200)
        return load_workbook(io.BytesIO(r.content))

    @staticmethod
    def sheet_values(wb, title_part):
        name = next(n for n in wb.sheetnames if title_part in n)
        return [list(r) for r in wb[name].iter_rows(values_only=True)]


class AccessTests(Base):
    def test_anonymous_goes_to_login(self):
        for u in (INDEX, export("xlsx"), export("pdf")):
            r = Client().get(u)
            self.assertEqual(r.status_code, 302)
            self.assertIn("/accounts/login/", r["Location"])

    def test_roles_without_the_permission_get_403_everywhere(self):
        sysadm = make_user("sys", role="sysadmin")
        for user in (self.emp, sysadm):
            c = self.client_for(user)
            self.assertEqual(c.get(INDEX).status_code, 403, user.username)
            self.assertEqual(c.get(export("xlsx"), self.params()).status_code, 403)
            self.assertEqual(c.get(export("pdf"), self.params()).status_code, 403)
        self.assertNotContains(self.client_for(self.emp).get(reverse("portal:dashboard")), INDEX)

    def test_view_without_export_permission_cannot_export(self):
        from django.contrib.auth.models import Permission

        g = Group.objects.get(name="hr")
        g.permissions.remove(Permission.objects.get(codename="export_reports"))
        for fmt in ("xlsx", "pdf"):
            self.assertEqual(self.client_for(self.hr).get(export(fmt), self.params()).status_code, 403)
        self.assertNotContains(self.page(self.hr), "data-export")

    def test_allowed_roles(self):
        for user in (self.admin, self.hr, self.head):
            self.assertEqual(self.page(user).status_code, 200, user.username)

    def test_export_accepts_only_get_and_known_formats(self):
        c = self.client_for(self.admin)
        self.assertEqual(c.post(export("xlsx"), self.params()).status_code, 405)
        self.assertEqual(c.get(export("csv"), self.params()).status_code, 404)
        self.assertEqual(c.get("/portal/reports/export/..%2Fetc/", self.params()).status_code, 404)

    def test_superuser_sees_everything(self):
        from django.contrib.auth import get_user_model

        su = get_user_model().objects.create_superuser("root", "root@example.test", "Namuna-Parol-12345")
        r = self.page(su)
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Moliya")


class FigureTests(Base):
    """Exact numbers from the known fixtures."""

    def test_admin_figures(self):
        d = self.page(self.admin).context["data"]
        k = d.kpis
        self.assertEqual(k["employees_active"], 3)  # a1 a2 b1 ; on_leave is not "active", archived excluded
        self.assertEqual(k["employees_total"], 4)  # + the one on leave
        self.assertEqual(k["tasks_total"], 5)  # the ancient task is outside the period
        self.assertEqual(k["tasks_done"], 3)
        self.assertEqual(k["tasks_done_rate"], 60)
        self.assertEqual(k["tasks_overdue_open"], 1)
        self.assertEqual(k["tasks_done_late"], 1)  # only "done late"
        self.assertEqual(k["tasks_completed_in_period"], 3)
        self.assertEqual(dict(d.table("task_status").rows), {"Yangi": 1, "Jarayonda": 1, "Tekshiruvda": 0, "Bajarildi": 3, "Bekor qilindi": 0})

    def test_departments_positions_documents(self):
        d = self.page(self.admin).context["data"]
        self.assertEqual(dict(d.table("emp_dept").rows), {"Maktabgacha": 2, "Moliya": 2})
        self.assertEqual(dict(d.table("emp_pos").rows), {"Metodist": 2, "Buxgalter": 2})
        self.assertEqual(dict(d.table("doc_cat").rows), {"Buyruqlar": 3, "Hisobotlar": 2})  # admin may open all five
        self.assertEqual(d.kpis["documents_total"], 5)

    def test_performance_tables(self):
        d = self.page(self.admin).context["data"]
        by_dept = {r[0]: r[1:] for r in d.table("perf_dept").rows}
        self.assertEqual(by_dept["Maktabgacha"], (4, 2, 1, "50%"))  # done on time, done late, overdue, fresh
        self.assertEqual(by_dept["Moliya"], (1, 1, 0, "100%"))
        by_user = {r[0]: r[1:] for r in d.table("perf_user").rows}
        self.assertEqual(by_user["Aliyev Vali"], (3, 2, 0, "67%"))
        self.assertEqual(by_user["Sodiqova Gulnora"], (2, 1, 1, "50%"))
        self.assertEqual(by_user["Botirov Sher"], (1, 1, 0, "100%"))

    def test_timeline_sums_match_the_kpis(self):
        d = self.page(self.admin).context["data"]
        self.assertEqual(sum(c for _, c, _ in d.timeline), 5)
        self.assertEqual(sum(x for _, _, x in d.timeline), 3)
        self.assertEqual(len(d.timeline), 31)  # one bucket per day of the 31-day window
        by_label = {l: (c, x) for l, c, x in d.timeline}
        day = (self.today - datetime.timedelta(days=10)).strftime("%d.%m")
        self.assertEqual(by_label[day], (1, 0))

    def test_week_and_month_grouping(self):
        for period, expected in (("week", None), ("month", None)):
            d = self.page(self.admin, period=period).context["data"]
            self.assertEqual(d.period, period)
            self.assertEqual(sum(c for _, c, _ in d.timeline), 5)
            self.assertEqual(sum(x for _, _, x in d.timeline), 3)

    def test_automatic_grouping_depends_on_length(self):
        long = self.page(self.admin, date_from=(self.today - datetime.timedelta(days=300)).isoformat()).context["data"]
        self.assertEqual(long.period, "month")
        self.assertEqual(sum(c for _, c, _ in long.timeline), 6)  # now the ancient task counts too

    def test_page_shows_the_numbers(self):
        r = self.page(self.admin)
        self.assertContains(r, "Faol xodimlar")
        self.assertContains(r, "Aliyev Vali")


class FilterTests(Base):
    def test_status_filter(self):
        d = self.page(self.admin, status="done").context["data"]
        self.assertEqual(d.kpis["tasks_total"], 3)
        self.assertEqual(d.kpis["tasks_overdue_open"], 0)

    def test_department_filter_includes_sub_departments_and_narrows_everything(self):
        d = self.page(self.admin, department=str(self.root.pk)).context["data"]
        self.assertEqual(d.kpis["tasks_total"], 4)
        self.assertEqual(dict(d.table("emp_dept").rows), {"Maktabgacha": 2})
        self.assertNotIn("Botirov Sher", [r[0] for r in d.table("perf_user").rows])

    def test_date_filter(self):
        narrow = self.page(self.admin, date_from=(self.today - datetime.timedelta(days=2)).isoformat()).context["data"]
        self.assertEqual(narrow.kpis["tasks_total"], 1)  # only "fresh"
        empty = self.page(self.admin, date_from=(self.today - datetime.timedelta(days=400)).isoformat(), date_to=(self.today - datetime.timedelta(days=380)).isoformat()).context["data"]
        self.assertEqual(empty.kpis["tasks_total"], 0)

    def test_category_filter(self):
        d = self.page(self.admin, category=str(self.cat2.pk)).context["data"]
        self.assertEqual(dict(d.table("doc_cat").rows), {"Hisobotlar": 2})

    def test_invalid_filters_show_errors_and_no_data(self):
        for bad, msg in (
            ({"date_from": self.today.isoformat(), "date_to": (self.today - datetime.timedelta(days=3)).isoformat()}, "keyin bo‘lishi mumkin emas"),
            ({"date_from": (self.today - datetime.timedelta(days=500)).isoformat()}, "oshmasligi kerak"),
            ({"date_from": "not-a-date"}, "to‘g‘ri kiriting"),
            ({"department": "garbage"}, ""),
            ({"status": "weird"}, ""),
        ):
            r = self.client_for(self.admin).get(INDEX, self.params(**bad))
            self.assertEqual(r.status_code, 200, bad)
            self.assertNotIn("data", r.context, bad)
            if msg:
                self.assertContains(r, msg)
            x = self.client_for(self.admin).get(export("xlsx"), self.params(**bad))
            self.assertEqual(x.status_code, 302, bad)  # redirected back with a message, no file

    def test_defaults_are_the_last_30_days(self):
        r = self.client_for(self.admin).get(INDEX)
        f = r.context["data"].filters
        self.assertEqual((f["date_to"], (f["date_to"] - f["date_from"]).days), (self.today, 29))


class ScopeAndPrivacyTests(Base):
    def test_department_head_sees_only_the_branch(self):
        r = self.page(self.head)
        body = r.content.decode()
        d = r.context["data"]
        self.assertEqual(d.kpis["tasks_total"], 4)
        self.assertEqual(dict(d.table("emp_dept").rows), {"Maktabgacha": 2})
        for hidden in ("Botirov", "Moliya", "Buxgalter", "finance"):
            self.assertNotIn(hidden, body, hidden)

    def test_head_cannot_widen_scope_with_a_foreign_department_parameter(self):
        r = self.client_for(self.head).get(INDEX, self.params(department=str(self.other.pk)))
        self.assertEqual(r.status_code, 200)
        self.assertNotIn("data", r.context)  # invalid choice: nothing is computed
        self.assertNotContains(r, "Botirov")
        x = self.client_for(self.head).get(export("xlsx"), self.params(department=str(self.other.pk)))
        self.assertEqual(x.status_code, 302)

    def test_head_sees_the_branch_dropdown_only(self):
        body = self.page(self.head).content.decode()
        self.assertIn("Maktabgacha", body)
        self.assertNotIn(f'value="{self.other.pk}"', body)

    def test_people_table_needs_its_own_permission(self):
        hr = self.page(self.hr)
        self.assertNotIn("perf_user", [t.key for t in hr.context["data"].tables])
        self.assertNotContains(hr, "Aliyev Vali")
        self.assertNotContains(hr, "Sodiqova")
        wb = self.workbook(self.hr)
        self.assertFalse(any("ijrochilar" in n for n in wb.sheetnames))
        flat = " ".join(str(c) for ws in wb for row in ws.iter_rows(values_only=True) for c in row if c)
        self.assertNotIn("Aliyev", flat)
        self.assertNotIn("Sodiqova", flat)

    def test_head_people_table_limited_to_the_branch_in_page_and_excel(self):
        names = [r[0] for r in self.page(self.head).context["data"].table("perf_user").rows]
        self.assertEqual(sorted(names), ["Aliyev Vali", "Sodiqova Gulnora"])
        wb = self.workbook(self.head)
        flat = " ".join(str(c) for ws in wb for row in ws.iter_rows(values_only=True) for c in row if c)
        self.assertNotIn("Botirov", flat)
        self.assertNotIn("Moliya", flat)
        self.assertNotIn("Buxgalter", flat)

    def test_documents_counts_never_include_documents_the_user_cannot_open(self):
        # the head holds documents.view_document but not view_all/confidential: general + own-branch departments only
        d = self.page(self.head).context["data"]
        self.assertEqual(dict(d.table("doc_cat").rows), {"Buyruqlar": 2, "Hisobotlar": 1})
        self.assertEqual(d.kpis["documents_total"], 3)
        wb = self.workbook(self.head)
        rows = self.sheet_values(wb, "turkumlar bo‘yicha")
        self.assertEqual(sorted((r[0], r[1]) for r in rows[1:] if r[0] and r[1] is not None and isinstance(r[1], int)), [("Buyruqlar", 2), ("Hisobotlar", 1)])

    def test_hr_documents_include_confidential_because_hr_may_open_them(self):
        self.assertEqual(self.page(self.hr).context["data"].kpis["documents_total"], 5)

    def test_users_without_document_access_get_no_document_section(self):
        from django.contrib.auth.models import Permission

        for name in ("department_head", "employee"):  # the head holds both roles
            Group.objects.get(name=name).permissions.remove(Permission.objects.get(codename="view_document"))
        r = self.page(self.head)
        self.assertNotIn("doc_cat", [t.key for t in r.context["data"].tables])
        wb = self.workbook(self.head)
        self.assertFalse(any("Hujjatlar" in n for n in wb.sheetnames))

    def test_head_without_a_branch_has_an_empty_scope(self):
        loose = make_user("loose", role="employee")
        RoleAssignment.objects.create(user=loose, role=Group.objects.get(name="department_head"), scope="department", department=Department.objects.create(name="Bo‘sh"))
        d = self.page(loose).context["data"]
        self.assertEqual((d.kpis["employees_total"], d.kpis["tasks_total"]), (0, 0))

    def test_build_report_refuses_users_without_permission_and_foreign_departments(self):
        from .services import FilterForm, build_report

        cleaned = {"date_from": self.today, "date_to": self.today, "department": None, "status": "", "category": None, "period": ""}
        with self.assertRaises(PermissionError):
            build_report(self.emp, cleaned)
        with self.assertRaises(PermissionError):
            build_report(self.head, {**cleaned, "department": self.other})


class ExportTests(Base):
    def test_xlsx_matches_the_page(self):
        wb = self.workbook(self.admin)
        self.assertIn("Umumiy", wb.sheetnames)
        summary = {r[0]: r[1] for r in self.sheet_values(wb, "Umumiy") if len(r) > 1 and r[1] is not None}
        self.assertEqual(summary["Faol xodimlar"], 3)
        self.assertEqual(summary["Davrda yaratilgan topshiriqlar"], 5)
        self.assertEqual(summary["Muddati o‘tgan (ochiq)"], 1)
        perf = {r[0]: r[1:4] for r in self.sheet_values(wb, "ijrochilar")[1:] if r[0]}
        self.assertEqual(tuple(perf["Aliyev Vali"]), (3, 2, 0))
        dyn = self.sheet_values(wb, "Dinamika")
        self.assertEqual((sum(r[1] for r in dyn[1:]), sum(r[2] for r in dyn[1:])), (5, 3))

    def test_xlsx_respects_filters(self):
        wb = self.workbook(self.admin, department=str(self.root.pk), status="done")
        summary = {r[0]: r[1] for r in self.sheet_values(wb, "Umumiy") if len(r) > 1 and r[1] is not None}
        self.assertEqual(summary["Davrda yaratilgan topshiriqlar"], 2)
        flat = " ".join(str(c) for ws in wb for row in ws.iter_rows(values_only=True) for c in row if c)
        self.assertNotIn("Botirov", flat)

    def test_headers_and_names(self):
        r = self.client_for(self.admin).get(export("xlsx"), self.params())
        self.assertEqual(r["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertRegex(r["Content-Disposition"], r'^attachment; filename="hisobot-\d{8}-\d{4}\.xlsx"$')
        self.assertIn("no-store", r["Cache-Control"])
        self.assertEqual(r["X-Content-Type-Options"], "nosniff")
        p = self.client_for(self.admin).get(export("pdf"), self.params())
        self.assertEqual(p["Content-Type"], "application/pdf")
        self.assertRegex(p["Content-Disposition"], r'filename="hisobot-\d{8}-\d{4}\.pdf"$')

    def test_pdf_is_a_valid_pdf_with_an_embedded_unicode_font(self):
        r = self.client_for(self.admin).get(export("pdf"), self.params())
        self.assertTrue(r.content.startswith(b"%PDF-"))
        self.assertIn(b"%%EOF", r.content[-64:])
        self.assertIn(b"DejaVu", r.content)
        self.assertGreater(len(r.content), 5000)

    def test_cyrillic_and_uzbek_characters_survive(self):
        u = make_user("uz1", role="employee")
        u.last_name, u.first_name = "Ғуломов", "Ўткир"
        u.save()
        Employee.objects.create(last_name="Ғуломов", first_name="Ўткир", department=self.child, position=Position.objects.create(name="Mutaxassis ‘oliy’ toifa ʻilm"), user=u, status="active")
        t = Task.objects.create(title="Hisobot", created_by=self.admin, status="new")
        t.assignees.set([u])
        wb = self.workbook(self.admin)
        flat = " ".join(str(c) for ws in wb for row in ws.iter_rows(values_only=True) for c in row if c)
        self.assertIn("Ғуломов Ўткир", flat)
        self.assertIn("Mutaxassis ‘oliy’ toifa ʻilm", flat)
        r = self.client_for(self.admin).get(export("pdf"), self.params())
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content.startswith(b"%PDF-"))

    def test_spreadsheet_formula_injection_is_neutralised(self):
        Department.objects.create(name="=HYPERLINK(\"http://evil.test\",\"x\")")
        Position.objects.create(name="@SUM(1+1)")
        u = make_user("inj", role="employee")
        u.last_name, u.first_name = "+cmd|' /C calc'!A0", "-2+3"
        u.save()
        Employee.objects.create(last_name="q", first_name="q", department=Department.objects.get(name__startswith="="), position=Position.objects.get(name="@SUM(1+1)"), user=u, status="active")
        t = Task.objects.create(title="x", created_by=self.admin)
        t.assignees.set([u])
        wb = self.workbook(self.admin)
        for ws in wb:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        self.assertFalse(cell.value.startswith(("=", "+", "-", "@")), (ws.title, cell.value))
                    self.assertNotEqual(cell.data_type, "f", (ws.title, cell.value))
        flat = " ".join(str(c) for ws in wb for row in ws.iter_rows(values_only=True) for c in row if c)
        self.assertIn("'=HYPERLINK", flat)

    def test_empty_result_still_produces_valid_files(self):
        far = {"date_from": (self.today - datetime.timedelta(days=400)).isoformat(), "date_to": (self.today - datetime.timedelta(days=390)).isoformat()}
        wb = self.workbook(self.admin, **far)
        summary = {r[0]: r[1] for r in self.sheet_values(wb, "Umumiy") if len(r) > 1 and r[1] is not None}
        self.assertEqual(summary["Davrda yaratilgan topshiriqlar"], 0)
        self.assertEqual(summary["Bajarilish ulushi, %"], 0)
        p = self.client_for(self.admin).get(export("pdf"), self.params(**far))
        self.assertTrue(p.content.startswith(b"%PDF-"))
        page = self.client_for(self.admin).get(INDEX, self.params(**far))
        self.assertContains(page, "ma’lumot yo‘q")

    def test_user_with_no_data_at_all(self):
        for m in (Task, Employee, Document):
            m.objects.all().delete()
        self.assertEqual(self.client_for(self.admin).get(export("xlsx"), self.params()).status_code, 200)
        self.assertEqual(self.client_for(self.admin).get(export("pdf"), self.params()).status_code, 200)
        self.assertEqual(self.page(self.admin).status_code, 200)

    @override_settings(REPORT_MAX_ROWS=3, REPORT_PDF_MAX_ROWS=2)
    def test_row_caps_protect_against_unbounded_output(self):
        for i in range(12):
            Department.objects.create(name=f"Bo‘lim {i:02d}")
            Employee.objects.create(last_name=f"E{i}", first_name="x", department=Department.objects.get(name=f"Bo‘lim {i:02d}"), status="active")
        page = self.page(self.admin)
        t = page.context["data"].table("emp_dept")
        self.assertEqual((len(t.rows), t.total_rows > 3), (3, True))
        self.assertContains(page, "Birinchi 3 ta qator")
        wb = self.workbook(self.admin)
        rows = self.sheet_values(wb, "bo‘limlar bo‘yicha")
        self.assertTrue(any(r[0] and "ko‘rsatilmadi" in str(r[0]) for r in rows))
        self.assertLessEqual(len(rows), 3 + 1 + 1 + 1)
        pdf = self.client_for(self.admin).get(export("pdf"), self.params())
        self.assertEqual(pdf.status_code, 200)

    def test_exports_are_audited_without_data(self):
        self.client_for(self.admin).get(export("xlsx"), self.params(department=str(self.root.pk)))
        e = AuditLog.objects.get(action=AuditLog.Action.REPORT_EXPORTED)
        self.assertEqual(e.user, self.admin)
        self.assertEqual(e.meta["format"], "xlsx")
        self.assertEqual(e.meta["filters"]["department"], str(self.root.pk))
        self.assertIn("people", e.meta["sections"])
        blob = str(e.meta)
        for name in ("Aliyev", "Botirov", "Sodiqova"):
            self.assertNotIn(name, blob)

    def test_denied_and_invalid_exports_are_not_audited_as_exports(self):
        self.client_for(self.emp).get(export("xlsx"), self.params())
        self.client_for(self.admin).get(export("xlsx"), self.params(date_from="x"))
        self.assertFalse(AuditLog.objects.filter(action=AuditLog.Action.REPORT_EXPORTED).exists())

    def test_viewing_the_people_table_is_audited_but_not_the_aggregates_only_view(self):
        self.page(self.hr)
        self.assertFalse(AuditLog.objects.filter(action=AuditLog.Action.REPORT_PEOPLE_VIEWED).exists())
        self.page(self.head)
        self.assertEqual(AuditLog.objects.filter(action=AuditLog.Action.REPORT_PEOPLE_VIEWED, user=self.head).count(), 1)


class RobustnessTests(Base):
    def test_a_failing_query_shows_an_error_state_not_a_traceback(self):
        from unittest import mock

        with mock.patch("apps.reports.views.build_report", side_effect=RuntimeError("secret internals")):
            r = self.client_for(self.admin).get(INDEX, self.params())
        self.assertEqual(r.status_code, 500)
        self.assertContains(r, "tayyorlab bo‘lmadi", status_code=500)
        self.assertNotContains(r, "secret internals", status_code=500)

    def test_long_period_is_capped_by_the_form(self):
        r = self.client_for(self.admin).get(INDEX, self.params(date_from=(self.today - datetime.timedelta(days=366)).isoformat()))
        self.assertNotIn("data", r.context)
        ok = self.client_for(self.admin).get(INDEX, self.params(date_from=(self.today - datetime.timedelta(days=365)).isoformat()))
        self.assertIn("data", ok.context)

    def test_html_in_names_is_escaped_on_the_page(self):
        Department.objects.create(name="<img src=x onerror=alert(1)>")
        Employee.objects.create(last_name="q", first_name="q", department=Department.objects.get(name__startswith="<img"), status="active")
        self.assertNotContains(self.page(self.admin), "<img src=x")

    def test_no_inline_script_on_the_page(self):
        body = self.page(self.admin).content.decode()
        self.assertFalse(re.search(r"<script(?![^>]*\bsrc=)", body))
