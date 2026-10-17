"""Employee import from Excel: upload -> preview -> confirm -> report. Fictional data only."""
import datetime
import io
import zipfile
from unittest import mock

from django.contrib.auth.models import Group
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from apps.access.models import RoleAssignment
from apps.accounts.testing import login_with_two_factor, make_user
from apps.audit.models import AuditLog

from . import importer
from .models import Department, Employee, EmployeeImport, Position

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def url(name, *args):
    return reverse(f"employees:{name}", args=args)


def xlsx(rows, header=None, *, extra_sheets=()):
    wb = Workbook()
    ws = wb.active
    ws.append(header if header is not None else importer.COLUMNS)
    for r in rows:
        ws.append(r)
    for title in extra_sheets:
        wb.create_sheet(title)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def row(last="Sinov", first="Rasul", middle="", dept="Bo‘lim A", pos="", status="", wphone="", wmail="", hired="", pphone="", birth=""):
    return [last, first, middle, dept, pos, status, wphone, wmail, hired, pphone, birth]


def up(data, name="xodimlar.xlsx"):
    return SimpleUploadedFile(name, data, content_type="")


class Base(TestCase):
    def setUp(self):
        call_command("seed_roles", verbosity=0)
        self.root = Department.objects.create(name="Bo‘lim A")
        self.child = Department.objects.create(name="Guruh A1", parent=self.root)
        self.b = Department.objects.create(name="Bo‘lim B")
        self.b_child = Department.objects.create(name="Guruh", parent=self.b)
        self.a_child = Department.objects.create(name="Guruh", parent=self.root)  # same name under another parent
        self.pos = Position.objects.create(name="Metodist")
        self.hr, self.admin = make_user("hrista", role="hr"), make_user("boss", role="admin")
        self.emp, self.stranger = make_user("plain", role="employee"), make_user("stranger", role="employee")
        self.content, self.sysadm = make_user("content", role="content_manager"), make_user("sysadm", role="sysadmin")
        self.head = make_user("head", role="employee")
        RoleAssignment.objects.create(user=self.head, role=Group.objects.get(name="department_head"), scope="department", department=self.root)
        self.c = self.client_for(self.hr)

    def client_for(self, user, csrf=False):
        c = Client(enforce_csrf_checks=csrf)
        login_with_two_factor(c, user)
        return c

    def upload(self, data, client=None, name="xodimlar.xlsx"):
        return (client or self.c).post(url("import"), {"file": up(data, name)})

    def preview(self, rows, header=None, **kw):
        r = self.upload(xlsx(rows, header), **kw)
        self.assertEqual(r.status_code, 302, r.content.decode()[:300] if r.status_code != 302 else "")
        return EmployeeImport.objects.latest("created_at")

    def confirm(self, imp, client=None, **data):
        return (client or self.c).post(url("import_confirm", imp.pk), data)

    def actions(self):
        return list(AuditLog.objects.values_list("action", flat=True))


class HappyPath(Base):
    def test_upload_previews_without_saving_then_confirm_creates_exactly_the_valid_rows(self):
        before = Employee.objects.count()
        imp = self.preview([
            row("Aliyev", "Vali", "Karimovich", "Bo‘lim A", "Metodist", "active", "+998 71 111 22 33", "vali@work.example.test", "2020-09-01", "+998 90 123 45 67", "1990-05-17"),
            row("Botirova", "Dilnoza", dept="Bo‘lim A / Guruh A1", status="faol", hired="01.02.2021"),
            row("Karimov", "Sardor", dept="Bo‘lim B"),
        ])
        self.assertEqual(Employee.objects.count(), before)  # a preview saves nothing
        page = self.c.get(url("import_detail", imp.pk))
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "Aliyev")
        self.assertEqual(page.context["tally"], {"ok": 3, "duplicate": 0, "error": 0, "total": 3})
        r = self.confirm(imp)
        self.assertRedirects(r, url("import_detail", imp.pk), fetch_redirect_response=False)
        self.assertEqual(Employee.objects.count(), before + 3)
        e = Employee.objects.get(last_name="Aliyev")
        self.assertEqual((e.first_name, e.middle_name, e.department, e.position, e.status), ("Vali", "Karimovich", self.root, self.pos, "active"))
        self.assertEqual((e.work_phone, e.work_email, e.hired_at, e.personal_phone, e.birth_date), ("+998 71 111 22 33", "vali@work.example.test", datetime.date(2020, 9, 1), "+998 90 123 45 67", datetime.date(1990, 5, 17)))
        d = Employee.objects.get(last_name="Botirova")
        self.assertEqual((d.department, d.status, d.hired_at, d.position), (self.child, "active", datetime.date(2021, 2, 1), None))  # localized status label + path + DD.MM.YYYY
        self.assertEqual(Employee.objects.get(last_name="Karimov").status, "new")  # blank status = new
        imp.refresh_from_db()
        self.assertEqual((imp.status, imp.rows, imp.summary["added"], imp.summary["total"]), ("completed", [], 3, 3))
        self.assertIsNotNone(imp.completed_at)
        report = self.c.get(url("import_detail", imp.pk))
        self.assertEqual(report.status_code, 200)
        self.assertNotContains(report, "Aliyev")  # the report carries numbers, not people

    def test_confirming_twice_creates_nothing_the_second_time(self):
        imp = self.preview([row("Aliyev", "Vali")])
        self.confirm(imp)
        n = Employee.objects.count()
        self.confirm(imp)
        self.assertEqual(Employee.objects.count(), n)
        self.assertEqual(self.actions().count("employee_import_completed"), 1)

    def test_the_downloadable_template_is_a_valid_file_and_imports_as_is(self):
        r = self.c.get(url("import_template"))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["X-Content-Type-Options"], "nosniff")
        self.assertIn("attachment", r["Content-Disposition"])
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual([c.value for c in wb["employees"][1]], importer.COLUMNS)
        self.assertIn("instructions", wb.sheetnames)
        text = " ".join(str(c.value) for rw in wb["instructions"].iter_rows() for c in rw)
        for word in ("O‘ZBEKCHA", "РУССКИЙ", "ENGLISH"):
            self.assertIn(word, text)
        Department.objects.create(name="Metodika bo‘limi")
        Department.objects.create(name="Maktabgacha guruh", parent=Department.objects.get(name="Metodika bo‘limi"))
        Position.objects.get_or_create(name="Metodist")
        imp = self.preview_file(r.content)
        self.assertEqual(self.c.get(url("import_detail", imp.pk)).context["tally"]["ok"], 2)

    def preview_file(self, data):
        self.assertEqual(self.upload(data).status_code, 302)
        return EmployeeImport.objects.latest("created_at")

    def test_blank_rows_are_skipped_and_excel_row_numbers_are_kept(self):
        imp = self.preview([row("Aliyev", "Vali"), [None] * 11, row("", "Bekzod")])
        results = importer.evaluate([(r["n"], r["raw"]) for r in imp.rows])
        self.assertEqual([(r["n"], r["state"]) for r in results], [(2, "ok"), (4, "error")])  # row 3 is blank, row 4 is the bad one

    def test_extra_sheets_are_ignored_and_only_the_first_is_read(self):
        imp = self.preview_file(xlsx([row("Aliyev", "Vali")], extra_sheets=("Boshqa",)))
        self.assertEqual(len(imp.rows), 1)


class RejectedFiles(Base):
    def rejected(self, data, name="x.xlsx", reason=None):
        n, open_before = Employee.objects.count(), EmployeeImport.objects.count()
        r = self.upload(data, name=name)
        self.assertEqual(r.status_code, 200)  # the upload page again, with the reason
        self.assertContains(r, "alert err")
        self.assertEqual((Employee.objects.count(), EmployeeImport.objects.count()), (n, open_before))  # no record, no rows kept
        if reason:
            self.assertIn(reason, [a.meta.get("reason") for a in AuditLog.objects.filter(action__in=["employee_import_rejected", "document_upload_rejected"])])
        return r

    def test_wrong_extension_wrong_content_and_missing_file(self):
        self.rejected(xlsx([row()]), name="x.csv", reason="not_xlsx")
        self.rejected(b"Familiya,Ism\nA,B\n", name="x.xlsx")
        self.rejected(b"%PDF-1.4\n%%EOF\n", name="x.xlsx")
        self.rejected(b"", name="x.xlsx")
        self.rejected(xlsx([row()]), name="x.ods")
        r = self.c.post(url("import"), {})
        self.assertContains(r, "alert err")

    def test_macro_and_active_content_files_are_refused(self):
        data = xlsx([row()])
        src = zipfile.ZipFile(io.BytesIO(data))
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as z:
            for n in src.namelist():
                z.writestr(n, src.read(n))
            z.writestr("xl/vbaProject.bin", b"\x00fake")
        self.rejected(out.getvalue(), reason="office_active")

    def test_xml_with_entities_is_refused_before_parsing(self):
        data = xlsx([row()])
        src = zipfile.ZipFile(io.BytesIO(data))
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as z:
            for n in src.namelist():
                body = src.read(n)
                if n == "xl/sharedStrings.xml" or n == "xl/worksheets/sheet1.xml":
                    body = b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaa">]>' + body.split(b"?>", 1)[-1]
                z.writestr(n, body)
        self.rejected(out.getvalue(), reason="xml_entities")

    def test_size_and_row_limits(self):
        with override_settings(EMPLOYEE_IMPORT_MAX_BYTES=2000):
            self.rejected(xlsx([row(f"Fam{chr(97 + i)}", "Ism") for i in range(20)]), reason="too_big")
        with override_settings(EMPLOYEE_IMPORT_MAX_ROWS=3):
            self.rejected(xlsx([row("Aaa", "Bbb")] * 4), reason="too_many_rows")
            self.assertEqual(self.upload(xlsx([row("Aaa", "Bbb")] * 3)).status_code, 302)
        # an "empty" sheet that is huge is cut off, not iterated for ever
        big = xlsx([row("Aaa", "Bbb")] + [[None] * 11] * (importer.MAX_SCANNED_ROWS + 10))
        self.rejected(big, reason="too_many_rows")

    def test_columns_are_checked_strictly(self):
        self.rejected(xlsx([["A", "B"]], header=["last_name", "first_name"]), reason="missing_columns")
        self.rejected(xlsx([row()], header=importer.COLUMNS[:-1] + ["salary"]), reason="unknown_columns")
        self.rejected(xlsx([row()], header=importer.COLUMNS + ["last_name"]), reason="duplicate_columns")
        self.rejected(xlsx([], header=importer.COLUMNS), reason="no_rows")
        wb = Workbook()
        wb.active.append([None, None])
        out = io.BytesIO()
        wb.save(out)
        self.rejected(out.getvalue(), reason="no_header")
        # headers are forgiving about case and spaces
        imp = self.preview([["Aliyev", "Vali", "Bo‘lim A"]], header=["Last Name", " FIRST_NAME ", "Department"])
        self.assertEqual(len(imp.rows), 1)

    def test_data_under_a_header_less_column_is_refused(self):
        self.rejected(xlsx([row() + ["surprise"]], header=importer.COLUMNS), reason="unknown_columns")


class RowRules(Base):
    def states(self, rows, **kw):
        imp = self.preview(rows, **kw)
        return importer.evaluate([(r["n"], r["raw"]) for r in imp.rows]), imp

    def codes(self, r):
        return [c for c, _p in r["errors"]]

    def test_every_kind_of_bad_value_is_reported_with_its_row_and_a_reason(self):
        results, imp = self.states([
            row("", "Vali"),                                  # 2 required
            row("Aliyev", "Vali", dept="Yo‘q bo‘lim"),        # 3 unknown department
            row("Aliyev", "Vali", dept="Guruh"),              # 4 ambiguous
            row("Aliyev", "Vali", pos="Yo‘q lavozim"),        # 5 unknown position
            row("Aliyev", "Vali", status="archived"),         # 6 import cannot archive
            row("Aliyev", "Vali", hired="31.02.2020"),        # 7 impossible date
            row("Aliyev", "Vali", birth="2015-01-01"),        # 8 birth out of range
            row("Aliyev", "Vali", wphone="abc"),              # 9 phone
            row("Aliyev", "Vali", wmail="not-an-email"),      # 10 e-mail
            row("=SUM(A1)", "Vali"),                          # 11 formula
            row("Ali7ev", "Vali"),                            # 12 digits in a name
            row("A" * 101, "Vali"),                           # 13 too long
            row("<script>alert(1)</script>", "Vali"),         # 14 markup in a name
        ])
        self.assertEqual([r["n"] for r in results], list(range(2, 15)))
        want = ["required", "dept_unknown", "dept_ambiguous", "position_unknown", "status_unknown", "date_format", "date_range", "phone_format", "email_format", "formula", "name_format", "too_long", "name_format"]
        for r, code in zip(results, want):
            self.assertEqual((r["state"], code in self.codes(r)), ("error", True), (r["n"], self.codes(r)))
        page = self.c.get(url("import_detail", imp.pk))
        self.assertEqual(page.context["tally"]["error"], 13)
        html = page.content.decode()
        self.assertNotIn("<script>alert(1)", html)
        self.assertIn("Yo‘q bo‘lim", html)
        for n in (2, 14):
            self.assertIn(f'<td class="nw">{n}</td>', html)

    def test_departments_by_name_by_path_and_ambiguity(self):
        results, _ = self.states([row("Aa", "Aa", dept="Guruh A1"), row("Bb", "Bb", dept="guruh a1"), row("Cc", "Cc", dept="Bo‘lim B / Guruh"),
                                  row("Dd", "Dd", dept="Bo‘lim A / Guruh"), row("Ee", "Ee", dept="Guruh"), row("Ff", "Ff", dept="A / Guruh A1")])
        self.assertEqual([r["state"] for r in results], ["ok", "ok", "ok", "ok", "error", "error"])
        self.assertEqual(results[1]["department_id"], str(self.child.pk))  # case-insensitive
        self.assertEqual(results[0]["department_id"], str(self.child.pk))
        self.assertEqual(results[2]["department_id"], str(self.b_child.pk))
        self.assertEqual(results[3]["department_id"], str(self.a_child.pk))
        self.assertIn("dept_ambiguous", self.codes(results[4]))
        self.assertIn("dept_unknown", self.codes(results[5]))
        self.root.is_active = False  # inactive departments are not valid targets
        self.child.is_active = False
        self.child.save()
        results, _ = self.states([row("Aaa", "Bbb", dept="Guruh A1")])
        self.assertIn("dept_unknown", self.codes(results[0]))

    def test_unusual_but_valid_names_pass(self):
        results, _ = self.states([row("O‘ktamov", "Sher’ali"), row("Abdullayev-Karimov", "Anvar"), row("Иванова", "Анна-Мария", "Петровна"), row("D’Artagnan", "Jo")])
        self.assertEqual([r["state"] for r in results], ["ok"] * 4)

    def test_formulas_are_never_evaluated_or_stored(self):
        results, _ = self.states([row("=1+1", "Vali"), row("Aliyev", "@cmd"), row("Aliyev", "Vali", wphone="=1+1")])
        self.assertEqual([r["state"] for r in results], ["error"] * 3)
        self.confirm(EmployeeImport.objects.latest("created_at"))
        self.assertFalse(Employee.objects.filter(last_name__startswith="=").exists())


class Duplicates(Base):
    def setUp(self):
        super().setUp()
        self.existing = Employee.objects.create(last_name="Aliyev", first_name="Vali", middle_name="Karimovich", department=self.b, status="active", work_email="vali@work.example.test", work_phone="111")
        self.gone = Employee.objects.create(last_name="Arxivov", first_name="Asad", department=self.b, status="archived")

    def test_same_name_email_or_an_archived_record_is_a_possible_duplicate_and_is_skipped_by_default(self):
        imp = self.preview([
            row("Aliyev", "Vali", "Karimovich"),                 # 2 same name, different department: ambiguous -> flagged, never merged
            row("ALIYEV", "vali", "karimovich"),                 # 3 case
            row("Aliyev", "Vali", "Karimovich", wmail="x@work.example.test"),  # 4
            row("Yangi", "Odam", wmail="VALI@work.example.test"),  # 5 e-mail of an existing person
            row("Arxivov", "Asad"),                              # 6 archived person still counts
            row("Yangiyev", "Sher"), row("Yangiyev", "Sher"),    # 7, 8 twice in the file
            row("Aliyev", "Vali"),                               # 9 different middle name (blank vs filled) -> NOT the same person
        ])
        results = importer.evaluate([(r["n"], r["raw"]) for r in imp.rows])
        self.assertEqual([(r["n"], r["state"]) for r in results], [(2, "duplicate"), (3, "duplicate"), (4, "duplicate"), (5, "duplicate"), (6, "duplicate"), (7, "ok"), (8, "duplicate"), (9, "ok")])
        self.assertIn("dup_file", self.codes(results[6]))
        before = Employee.objects.count()
        self.confirm(imp)
        self.assertEqual(Employee.objects.count(), before + 2)  # rows 7 and 9 only
        self.existing.refresh_from_db()
        self.assertEqual((self.existing.department, self.existing.work_phone, self.existing.status), (self.b, "111", "active"))  # untouched
        imp.refresh_from_db()
        self.assertEqual(sorted(imp.summary["duplicates_skipped"]), [2, 3, 4, 5, 6, 8])
        self.assertEqual(Employee.objects.filter(last_name="Aliyev", first_name="Vali").count(), 2)  # the original and row 9

    def test_importing_duplicates_needs_the_explicit_box_and_always_creates_new_rows(self):
        imp = self.preview([row("Aliyev", "Vali", "Karimovich", dept="Bo‘lim A")])
        before = Employee.objects.count()
        self.confirm(imp, include_duplicates="1")
        self.assertEqual(Employee.objects.count(), before + 1)
        self.assertEqual(Employee.objects.filter(last_name="Aliyev", first_name="Vali", middle_name="Karimovich").count(), 2)
        self.existing.refresh_from_db()
        self.assertEqual(self.existing.department, self.b)  # still the old one, nothing merged or overwritten
        imp.refresh_from_db()
        self.assertEqual(imp.summary["duplicates_imported"], [2])

    def test_a_duplicate_that_appears_between_preview_and_confirm_is_caught_at_confirm(self):
        imp = self.preview([row("Yangi", "Odam")])
        Employee.objects.create(last_name="Yangi", first_name="Odam", department=self.root, status="active")  # someone added them meanwhile
        before = Employee.objects.count()
        self.confirm(imp)
        self.assertEqual(Employee.objects.count(), before)

    def codes(self, r):
        return [c for c, _p in r["errors"]]


class PartialAndAtomic(Base):
    def test_a_file_with_bad_rows_imports_only_the_good_ones_and_reports_the_rest(self):
        imp = self.preview([row("Aliyev", "Vali"), row("", "Yo‘q"), row("Botirov", "Sher", dept="Yo‘q"), row("Karimov", "Said")])
        page = self.c.get(url("import_detail", imp.pk))
        self.assertEqual(page.context["tally"], {"ok": 2, "duplicate": 0, "error": 2, "total": 4})
        self.confirm(imp)
        self.assertEqual(sorted(Employee.objects.values_list("last_name", flat=True)), ["Aliyev", "Karimov"])
        imp.refresh_from_db()
        self.assertEqual([e["n"] for e in imp.summary["errors"]], [3, 4])
        report = self.c.get(url("import_detail", imp.pk))
        self.assertEqual([e["n"] for e in report.context["errors"]], [3, 4])
        self.assertTrue(all(e["messages"] for e in report.context["errors"]))

    def test_a_department_deactivated_after_the_preview_turns_the_row_into_a_reported_error(self):
        imp = self.preview([row("Aliyev", "Vali", dept="Bo‘lim B"), row("Karimov", "Said", dept="Bo‘lim A")])
        self.b.is_active = False
        self.b.save()
        self.confirm(imp)
        self.assertEqual(list(Employee.objects.values_list("last_name", flat=True)), ["Karimov"])
        imp.refresh_from_db()
        self.assertEqual([e["n"] for e in imp.summary["errors"]], [2])

    def test_a_failure_in_the_middle_rolls_everything_back_and_keeps_the_preview(self):
        imp = self.preview([row("Aliyev", "Vali"), row("Karimov", "Said"), row("Botirov", "Sher")])
        real = __import__("apps.organization.import_views", fromlist=["record"]).record
        calls = {"n": 0}

        def flaky(action, **kw):
            if action == AuditLog.Action.EMPLOYEE_CREATED:
                calls["n"] += 1
                if calls["n"] == 2:
                    raise RuntimeError("disk on fire")
            return real(action, **kw)

        before = Employee.objects.count()
        audit_before = AuditLog.objects.filter(action="employee_created").count()
        with mock.patch("apps.organization.import_views.record", side_effect=flaky):
            r = self.confirm(imp)
        self.assertEqual(r.status_code, 302)
        self.assertEqual(Employee.objects.count(), before)  # not one of the three
        self.assertEqual(AuditLog.objects.filter(action="employee_created").count(), audit_before)
        imp.refresh_from_db()
        self.assertEqual((imp.status, len(imp.rows)), ("previewed", 3))  # can be retried or cancelled
        self.assertIn("employee_import_failed", self.actions())
        self.assertNotIn("employee_import_completed", self.actions())
        self.confirm(imp)  # the retry works
        self.assertEqual(Employee.objects.count(), before + 3)


class Lifecycle(Base):
    def test_cancel_wipes_the_rows_and_creates_nothing(self):
        imp = self.preview([row("Aliyev", "Vali")])
        n = Employee.objects.count()
        self.c.post(url("import_cancel", imp.pk))
        imp.refresh_from_db()
        self.assertEqual((imp.status, imp.rows), ("cancelled", []))
        self.assertEqual(Employee.objects.count(), n)
        self.confirm(imp)
        self.assertEqual(Employee.objects.count(), n)
        self.assertIn("employee_import_cancelled", self.actions())

    def test_an_unconfirmed_preview_expires_and_its_personal_data_is_wiped(self):
        imp = self.preview([row("Aliyev", "Vali", pphone="+998 90 123 45 67")])
        EmployeeImport.objects.filter(pk=imp.pk).update(created_at=timezone.now() - datetime.timedelta(minutes=61))
        n = Employee.objects.count()
        page = self.c.get(url("import_detail", imp.pk))
        self.assertEqual(page.status_code, 200)
        self.assertNotContains(page, "Aliyev")
        self.confirm(imp)
        self.assertEqual(Employee.objects.count(), n)  # too late, even before anything wiped it
        call_command("purge_employee_imports", verbosity=0)
        imp.refresh_from_db()
        self.assertEqual((imp.status, imp.rows), ("expired", []))
        fresh = self.preview([row("Karimov", "Said")])
        call_command("purge_employee_imports", verbosity=0)
        fresh.refresh_from_db()
        self.assertEqual(fresh.status, "previewed")  # a young one is left alone

    def test_the_uploaded_file_is_never_written_to_disk(self):
        from django.conf import settings
        import os

        def count():
            return sum(len(f) for _r, _d, f in os.walk(settings.PRIVATE_MEDIA_ROOT)) if os.path.isdir(settings.PRIVATE_MEDIA_ROOT) else 0

        before = count()
        imp = self.preview([row("Aliyev", "Vali")])
        self.confirm(imp)
        self.assertEqual(count(), before)

    def test_audit_has_counts_and_ids_but_no_personal_data(self):
        imp = self.preview([row("Aliyev", "Vali", pphone="+998 90 555 66 77", birth="1990-05-17", wmail="vali@work.example.test"), row("", "Xato"), row("Aliyev", "Vali")])
        self.confirm(imp)
        rows = AuditLog.objects.filter(action__in=["employee_import_previewed", "employee_import_completed", "employee_created"])
        self.assertEqual(sorted(a.action for a in rows), ["employee_created", "employee_import_completed", "employee_import_previewed"])
        for a in rows:
            blob = str(a.meta) + a.object_id
            for secret in ("Aliyev", "Vali", "+998", "1990", "@work", "Xato"):
                self.assertNotIn(secret, blob, a.action)
            self.assertEqual(a.user, self.hr)
        done = AuditLog.objects.get(action="employee_import_completed")
        self.assertEqual((done.meta["added"], done.meta["skipped"], done.meta["errors"]), (1, 2, 1))
        self.assertEqual(done.object_id, str(imp.pk))
        created = AuditLog.objects.get(action="employee_created")
        self.assertEqual(created.meta, {"via": "import", "import": str(imp.pk)})
        self.assertEqual(AuditLog.objects.get(action="employee_import_previewed").meta["total"], 3)

    def test_rejected_uploads_are_audited_with_a_reason_only(self):
        self.upload(xlsx([row()], header=importer.COLUMNS[:-1] + ["salary"]))
        a = AuditLog.objects.get(action="employee_import_rejected")
        self.assertEqual(a.meta, {"reason": "unknown_columns"})


class Permissions(Base):
    def routes(self, imp):
        return [("get", url("import")), ("post", url("import")), ("get", url("import_template")), ("get", url("import_detail", imp.pk)),
                ("post", url("import_confirm", imp.pk)), ("post", url("import_cancel", imp.pk))]

    def test_only_people_who_may_add_employees_get_in_and_nothing_is_created_for_the_others(self):
        imp = self.preview([row("Aliyev", "Vali")])
        n = Employee.objects.count()
        for user in (self.emp, self.stranger, self.head, self.content, self.sysadm):
            c = self.client_for(user)
            for method, u in self.routes(imp):
                body = {"file": up(xlsx([row("Hacker", "One")]))} if (method == "post" and u == url("import")) else {}
                self.assertEqual(getattr(c, method)(u, body).status_code, 403, (user.username, u))
        for method, u in self.routes(imp):
            r = getattr(Client(), method)(u)
            self.assertEqual(r.status_code, 302)
            self.assertIn("/accounts/login/", r["Location"])
        self.assertEqual(Employee.objects.count(), n)
        imp.refresh_from_db()
        self.assertEqual(imp.status, "previewed")
        self.assertFalse(Employee.objects.filter(last_name="Hacker").exists())

    def test_admin_may_import_too_but_only_hr_sees_its_own_preview(self):
        admin = self.client_for(self.admin)
        self.assertEqual(admin.get(url("import")).status_code, 200)
        imp = self.preview([row("Aliyev", "Vali")])
        # another person who may import still cannot open, confirm or cancel someone else's preview
        for method, u in ((("get", url("import_detail", imp.pk))), ("post", url("import_confirm", imp.pk)), ("post", url("import_cancel", imp.pk))):
            self.assertEqual(getattr(admin, method)(u).status_code, 404)
        self.assertFalse(Employee.objects.filter(last_name="Aliyev").exists())
        imp.refresh_from_db()
        self.assertEqual((imp.status, len(imp.rows)), ("previewed", 1))

    def test_csrf_and_methods(self):
        imp = self.preview([row("Aliyev", "Vali")])
        c = self.client_for(self.hr, csrf=True)
        for u in (url("import"), url("import_confirm", imp.pk), url("import_cancel", imp.pk)):
            self.assertEqual(c.post(u, {}).status_code, 403, u)
        self.assertFalse(Employee.objects.filter(last_name="Aliyev").exists())
        for u in (url("import_confirm", imp.pk), url("import_cancel", imp.pk)):
            self.assertEqual(self.c.get(u).status_code, 405)
        for m in ("put", "patch", "delete"):
            self.assertEqual(getattr(self.c, m)(url("import")).status_code, 405)
        self.assertEqual(self.c.get(url("import_detail", imp.pk) + "?only=problems").status_code, 200)

    def test_unknown_or_malformed_ids_are_404_not_500(self):
        for u in (url("import_detail", "11111111-2222-3333-4444-555555555555"), url("import_confirm", "11111111-2222-3333-4444-555555555555")):
            self.assertIn(self.c.post(u).status_code if "confirm" in u else self.c.get(u).status_code, (404,))
        self.assertEqual(self.c.get("/portal/employees/import/not-a-uuid/").status_code, 404)


class Languages(Base):
    def test_pages_and_row_messages_follow_the_language(self):
        imp = self.preview([row("", "Vali"), row("Aliyev", "Vali", dept="Yo‘q")])
        for lang, word, msg in (("ru", "Импорт", "Обязательное поле не заполнено."), ("en", "Import", "Required field is empty."), ("uz", "import", "Majburiy maydon bo‘sh.")):
            c = self.client_for(self.hr)
            c.cookies["django_language"] = lang
            page = c.get(url("import_detail", imp.pk))
            self.assertContains(page, msg)
            self.assertContains(page, f'<html lang="{lang}">')
            self.assertContains(c.get(url("import")), f'<html lang="{lang}">')
        c = self.client_for(self.hr)
        c.cookies["django_language"] = "ru"
        bad = c.post(url("import"), {"file": up(b"junk")})
        self.assertContains(bad, "alert err")
        self.assertNotContains(bad, "Fayl")  # no Uzbek leaks into the Russian error
