"""Employee import from .xlsx: parse -> validate -> (later, on confirmation) create. Create-only: existing employees are never changed.

The file is read once, in memory, and never stored. Nothing in it is executed or fetched: no macros (rejected by the shared
upload validation), no formulas (any cell starting with "=" is an error), no external links (openpyxl does not follow them and
files with DTDs / entities are refused before parsing). Every row goes through ONE pipeline (`evaluate`), used for the preview
and again on confirmation against the then-current database."""
import datetime
import io
import re
import unicodedata
import zipfile

from django.conf import settings
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.utils import translation
from django.utils.translation import gettext_lazy as _

from .models import Department, Employee, Position

COLUMNS = ["last_name", "first_name", "middle_name", "department", "position", "status", "work_phone", "work_email", "hired_at", "personal_phone", "birth_date"]
REQUIRED = ["last_name", "first_name", "department"]
DATE_COLUMNS = {"hired_at", "birth_date"}
NAME_COLUMNS = {"last_name", "first_name", "middle_name"}
PHONE_COLUMNS = {"work_phone", "personal_phone"}
MAX_SCANNED_ROWS = 5000        # blank rows included: an "empty" sheet of a million rows must not cost a million iterations
MAX_XML_PART_BYTES = 20 * 1024 * 1024
PHONE_RE = re.compile(r"^\+?[0-9][0-9 ()\-]{4,30}$")
APOSTROPHES = "‘’ʻʼ`´"
MESSAGES = {
    "required": _("Majburiy maydon bo‘sh."),
    "too_long": _("Qiymat juda uzun (ko‘pi bilan %(max)d belgi)."),
    "bad_chars": _("Qiymatda ruxsat etilmagan belgilar bor."),
    "formula": _("Formula yoki “=”, “@” bilan boshlanadigan qiymat qabul qilinmaydi."),
    "name_format": _("Ism-familiya faqat harflar, bo‘sh joy, defis va apostrofdan iborat bo‘lishi kerak."),
    "dept_unknown": _("Bunday faol bo‘lim topilmadi: “%(value)s”. Import yangi bo‘lim yaratmaydi."),
    "dept_ambiguous": _("Bo‘lim nomi bir necha joyda uchraydi: “%(value)s”. To‘liq yo‘lni yozing (masalan, “Yuqori bo‘lim / Quyi bo‘lim”)."),
    "position_unknown": _("Bunday lavozim topilmadi: “%(value)s”. Import yangi lavozim yaratmaydi."),
    "status_unknown": _("Noma’lum holat: “%(value)s”. Ruxsat etilgan: %(allowed)s."),
    "phone_format": _("Telefon raqami noto‘g‘ri formatda."),
    "email_format": _("E-pochta manzili noto‘g‘ri."),
    "date_format": _("Sana noto‘g‘ri. Format: YYYY-MM-DD yoki DD.MM.YYYY."),
    "date_range": _("Sana ruxsat etilgan oraliqdan tashqarida."),
    "dup_existing": _("Bunday F.I.Sh. bilan xodim allaqachon mavjud (mumkin bo‘lgan dublikat). Mavjud xodim o‘zgartirilmaydi."),
    "dup_email": _("Bu ish e-pochtasi boshqa xodimda bor (mumkin bo‘lgan dublikat)."),
    "dup_file": _("Faylda xuddi shunday F.I.Sh. avval ham uchragan (mumkin bo‘lgan dublikat)."),
}
FILE_ERRORS = {
    "not_xlsx": _("Bu .xlsx fayli emas yoki u buzilgan."),
    "too_big": _("Fayl hajmi %(mb)d MB dan oshmasligi kerak."),
    "xml_entities": _("Fayl tarkibi xavfli (DTD yoki entity)."),
    "xml_big": _("Fayl tarkibi juda katta."),
    "no_sheet": _("Faylda ma’lumotli varaq yo‘q."),
    "no_header": _("Birinchi qatorda ustun nomlari bo‘lishi kerak."),
    "missing_columns": _("Majburiy ustunlar yo‘q: %(columns)s."),
    "unknown_columns": _("Noma’lum ustunlar: %(columns)s. Shablondagi ustunlardan foydalaning."),
    "duplicate_columns": _("Ustun nomi takrorlangan: %(columns)s."),
    "too_many_rows": _("Qatorlar soni %(max)d dan oshmasligi kerak."),
    "no_rows": _("Faylda import qilinadigan qatorlar yo‘q."),
}


class ImportFileError(Exception):
    def __init__(self, code, **params):
        super().__init__(code)
        self.code, self.params = code, params

    @property
    def message(self):
        text = str(FILE_ERRORS[self.code])
        return text % self.params if self.params else text


def message(code, params=None):
    text = str(MESSAGES[code])
    if "%(" not in text:
        return text
    return text % params if params else re.sub(r"%\(\w+\)[sd]", "…", text)


# ---- reading the workbook ---------------------------------------------------------------------------------------

def _check_package(data: bytes):
    """Extra checks on top of apps.documents.validation: no DTD/entity declarations and no oversized XML part, because the
    spreadsheet library would otherwise expand them."""
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise ImportFileError("not_xlsx") from exc
    for info in zf.infolist():
        if not info.filename.lower().endswith((".xml", ".rels")):
            continue
        if info.file_size > MAX_XML_PART_BYTES:
            raise ImportFileError("xml_big")
        head = zf.read(info.filename)
        if b"<!DOCTYPE" in head or b"<!ENTITY" in head:
            raise ImportFileError("xml_entities")


def read_rows(data: bytes):
    """-> list of (row_number, {column: raw cell value}); raises ImportFileError. Row numbers are the Excel row numbers."""
    from openpyxl import load_workbook

    if len(data) > settings.EMPLOYEE_IMPORT_MAX_BYTES:
        raise ImportFileError("too_big", mb=settings.EMPLOYEE_IMPORT_MAX_BYTES // (1024 * 1024))
    _check_package(data)
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        ws = wb.worksheets[0]
    except Exception as exc:  # openpyxl raises many types for broken packages
        raise ImportFileError("not_xlsx") from exc
    header, out, scanned = None, [], 0
    try:
        for number, cells in enumerate(ws.iter_rows(values_only=True), start=1):
            scanned += 1
            if scanned > MAX_SCANNED_ROWS:
                raise ImportFileError("too_many_rows", max=settings.EMPLOYEE_IMPORT_MAX_ROWS)
            cells = list(cells)
            while cells and cells[-1] in (None, ""):
                cells.pop()
            if not cells:
                continue
            if header is None:
                header = [_header(c) for c in cells]
                _check_header(header)
                continue
            if len(cells) > len(header):
                extra = [c for c in cells[len(header):] if c not in (None, "")]
                if extra:
                    raise ImportFileError("unknown_columns", columns=_("(sarlavhasiz ustun)"))
            row = {name: (cells[i] if i < len(cells) else None) for i, name in enumerate(header) if name}
            if all(v in (None, "") or (isinstance(v, str) and not v.strip()) for v in row.values()):
                continue
            out.append((number, row))
            if len(out) > settings.EMPLOYEE_IMPORT_MAX_ROWS:
                raise ImportFileError("too_many_rows", max=settings.EMPLOYEE_IMPORT_MAX_ROWS)
    except ImportFileError:
        raise
    except Exception as exc:
        raise ImportFileError("not_xlsx") from exc
    finally:
        wb.close()
    if header is None:
        raise ImportFileError("no_header")
    if not out:
        raise ImportFileError("no_rows")
    return out


def _header(cell):
    return re.sub(r"\s+", "_", str(cell).strip().lower()) if cell not in (None, "") else ""


def _check_header(header):
    named = [h for h in header if h]
    dupes = sorted({h for h in named if named.count(h) > 1})
    if dupes:
        raise ImportFileError("duplicate_columns", columns=", ".join(dupes))
    unknown = [h for h in named if h not in COLUMNS]
    if unknown:
        raise ImportFileError("unknown_columns", columns=", ".join(unknown)[:200])
    missing = [c for c in REQUIRED if c not in named]
    if missing:
        raise ImportFileError("missing_columns", columns=", ".join(missing))


# ---- cleaning and validating one row ----------------------------------------------------------------------------

def _text(value):
    """Cell -> clean string (NFC, one-space, no control characters); raises ValueError('bad_chars') or ('formula')."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.date().isoformat() if isinstance(value, datetime.datetime) else value.isoformat()
    text = unicodedata.normalize("NFC", str(value))
    if any(unicodedata.category(ch) in ("Cc", "Cf", "Cs", "Co", "Cn") and ch not in "\t\n" for ch in text):
        raise ValueError("bad_chars")
    text = re.sub(r"\s+", " ", text).strip()
    if text.startswith(("=", "@")):
        raise ValueError("formula")
    return text


def clean_row(raw: dict):
    """-> (cleaned {column: str}, errors [(code, params)])."""
    cleaned, errors = {}, []
    for col in COLUMNS:
        try:
            cleaned[col] = _text(raw.get(col))
        except ValueError as exc:
            cleaned[col] = ""
            errors.append((str(exc), {"column": col}))
    return cleaned, errors


def norm_name(s):
    for a in APOSTROPHES:
        s = s.replace(a, "'")
    return re.sub(r"\s+", " ", s).casefold().strip()


def _name_key(c):
    return (norm_name(c["last_name"]), norm_name(c["first_name"]), norm_name(c["middle_name"]))


def _status_map():
    out = {"": Employee.Status.NEW}
    for lang, _name in settings.LANGUAGES:
        with translation.override(lang):
            for code, label in Employee.Status.choices:
                if code != Employee.Status.ARCHIVED:
                    out[norm_name(str(label))] = code
    for code in Employee.Status.values:
        if code != Employee.Status.ARCHIVED:
            out[code] = code
    return out


def _parse_date(text):
    for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


class Lookups:
    """Departments, positions, names and e-mails as they are NOW; built once per evaluation."""

    def __init__(self):
        self.departments = list(Department.objects.filter(is_active=True).values("id", "name", "parent_id"))
        self.by_id = {d["id"]: d for d in Department.objects.values("id", "name", "parent_id")}
        self.positions = {norm_name(p.name): p for p in Position.objects.all()}
        self.statuses = _status_map()
        self.names, self.emails = set(), set()
        for last, first, middle, email in Employee.objects.values_list("last_name", "first_name", "middle_name", "work_email"):
            self.names.add((norm_name(last), norm_name(first), norm_name(middle)))
            if email:
                self.emails.add(email.casefold())

    def department(self, value):
        parts = [norm_name(p) for p in value.split("/") if p.strip()]
        if not parts:
            return None, "dept_unknown"
        hits = []
        for d in self.departments:
            chain, cur, guard = [], d, 0
            while cur and guard < 50:  # root-first chain of names; the guard stops a (corrupt) cycle
                chain.insert(0, norm_name(cur["name"]))
                cur = self.by_id.get(cur["parent_id"])
                guard += 1
            if chain[-len(parts):] == parts:
                hits.append(d)
        if len(hits) == 1:
            return hits[0]["id"], None
        return None, "dept_unknown" if not hits else "dept_ambiguous"


def evaluate(parsed_rows, lookups=None):
    """parsed_rows: [(row_number, {column: raw})] -> list of result dicts:
    {"n", "state": ok|duplicate|error, "cleaned": {...}, "department_id", "position_id", "status", "errors": [[code, params]]}
    A row with errors is "error"; otherwise a name/e-mail match with an existing or earlier row makes it "duplicate"."""
    lk = lookups or Lookups()
    seen_names, seen_emails, results = set(), set(), []
    today = datetime.date.today()
    for number, raw in parsed_rows:
        c, errors = clean_row(raw)
        bad_cols = {p["column"] for code, p in errors}
        for col in REQUIRED:
            if col not in bad_cols and not c[col]:
                errors.append(("required", {"column": col}))
        for col in COLUMNS:
            if col in bad_cols or not c[col]:
                continue
            limit = 200 if col in ("department", "position") else 254 if col == "work_email" else 100 if col in NAME_COLUMNS else 32
            if len(c[col]) > limit:
                errors.append(("too_long", {"column": col, "max": limit}))
        for col in NAME_COLUMNS:
            v = c[col]
            if v and col not in bad_cols and (any(ch.isdigit() for ch in v) or not re.fullmatch(r"[^\W\d_]+(?:[ '.\-" + APOSTROPHES + r"][^\W\d_]*)*\.?", v)):
                errors.append(("name_format", {"column": col}))
        for col in PHONE_COLUMNS:
            if c[col] and col not in bad_cols and not PHONE_RE.match(c[col]):
                errors.append(("phone_format", {"column": col}))
        if c["work_email"] and "work_email" not in bad_cols:
            try:
                validate_email(c["work_email"])
            except ValidationError:
                errors.append(("email_format", {"column": "work_email"}))
        dept_id = pos_id = None
        status = Employee.Status.NEW
        if c["department"] and "department" not in bad_cols:
            dept_id, err = lk.department(c["department"])
            if err:
                errors.append((err, {"column": "department", "value": c["department"][:80]}))
        if c["position"] and "position" not in bad_cols:
            pos = lk.positions.get(norm_name(c["position"]))
            if pos is None:
                errors.append(("position_unknown", {"column": "position", "value": c["position"][:80]}))
            else:
                pos_id = pos.pk
        if "status" not in bad_cols:
            status = lk.statuses.get(norm_name(c["status"]))
            if status is None:
                allowed = ", ".join(sorted(code for code in Employee.Status.values if code != "archived"))
                errors.append(("status_unknown", {"column": "status", "value": c["status"][:40], "allowed": allowed}))
        dates = {}
        for col in DATE_COLUMNS:
            if not c[col] or col in bad_cols:
                continue
            d = _parse_date(c[col])
            if d is None:
                errors.append(("date_format", {"column": col}))
                continue
            low, high = (datetime.date(1930, 1, 1), today - datetime.timedelta(days=14 * 365)) if col == "birth_date" else (datetime.date(1960, 1, 1), today)
            if not low <= d <= high:
                errors.append(("date_range", {"column": col}))
                continue
            dates[col] = d.isoformat()
        c.update({k: dates.get(k, "") for k in DATE_COLUMNS})
        state, dup = "ok", []
        if errors:
            state = "error"
        else:
            key = _name_key(c)
            email = c["work_email"].casefold()
            if key in lk.names:
                dup.append("dup_existing")
            if email and email in lk.emails:
                dup.append("dup_email")
            if key in seen_names:
                dup.append("dup_file")
            if email and email in seen_emails:
                dup.append("dup_file")
            if dup:
                state = "duplicate"
            seen_names.add(key)
            if email:
                seen_emails.add(email)
        results.append({
            "n": number, "state": state, "cleaned": c, "department_id": str(dept_id) if dept_id else None, "position_id": str(pos_id) if pos_id else None,
            "status": status if isinstance(status, str) else str(status), "errors": [[code, params] for code, params in errors + [(d, {}) for d in dup]],
        })
    return results


def render_errors(errors):
    """[[code, params], ...] -> list of messages in the CURRENT language."""
    return [message(code, params) for code, params in errors]


def tally(results):
    return {s: sum(1 for r in results if r["state"] == s) for s in ("ok", "duplicate", "error")} | {"total": len(results)}


def build_template() -> bytes:
    """The example workbook: header row, two clearly fictional rows, and an instruction sheet in the three interface languages."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "employees"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.append(["Namunaviy", "Anvar", "Karimovich", "Metodika bo‘limi", "Metodist", "active", "+998 71 000 00 00", "anvar.namunaviy@example.test", "2020-09-01", "", ""])
    ws.append(["Sinov", "Dilnoza", "", "Metodika bo‘limi / Maktabgacha guruh", "", "new", "", "", "", "", ""])
    for col, width in zip("ABCDEFGHIJK", (16, 14, 16, 36, 16, 10, 20, 30, 12, 18, 12)):
        ws.column_dimensions[col].width = width
    guide = wb.create_sheet("instructions")
    for line in INSTRUCTIONS:
        guide.append([line])
    guide.column_dimensions["A"].width = 120
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


INSTRUCTIONS = [
    "O‘ZBEKCHA — Birinchi varaq (employees): 1-qator ustun nomlari (o‘zgartirmang). Majburiy: last_name, first_name, department. Qolganlari ixtiyoriy.",
    "department: bo‘lim nomi, nom takrorlansa — to‘liq yo‘l “Yuqori / Quyi”. position: mavjud lavozim nomi. status: new (bo‘sh = new), active, on_leave.",
    "Sanalar: YYYY-MM-DD yoki DD.MM.YYYY. Telefon: +998 71 000 00 00. Formulalar, makroslar, havolalar qabul qilinmaydi. Ko‘pi bilan 500 qator, 2 MB.",
    "Import faqat YANGI xodimlar qo‘shadi; mavjudlarni o‘zgartirmaydi va birlashtirmaydi. F.I.Sh. yoki e-pochta mos kelsa — “mumkin bo‘lgan dublikat”, sukut bo‘yicha o‘tkazib yuboriladi.",
    "Bo‘lim va lavozim oldindan portalda yaratilgan bo‘lishi kerak. Namuna qatorlari to‘qima; ularni o‘chirib, o‘z ma’lumotlaringizni kiriting.",
    "РУССКИЙ — Первый лист (employees): строка 1 — названия столбцов (не меняйте). Обязательные: last_name, first_name, department. Остальные необязательны.",
    "department: название подразделения, а при совпадении названий — полный путь «Верхнее / Нижнее». position: название существующей должности. status: new (пусто = new), active, on_leave.",
    "Даты: YYYY-MM-DD или DD.MM.YYYY. Телефон: +998 71 000 00 00. Формулы, макросы и ссылки не принимаются. Не более 500 строк и 2 МБ.",
    "Импорт только ДОБАВЛЯЕТ новых сотрудников; существующих не изменяет и не объединяет. Совпадение ФИО или e-mail — «возможный дубликат», по умолчанию пропускается.",
    "Подразделения и должности должны быть заранее созданы в портале. Строки-примеры вымышлены: удалите их и внесите свои данные.",
    "ENGLISH — First sheet (employees): row 1 holds the column names (do not change them). Required: last_name, first_name, department. The rest are optional.",
    "department: department name, or the full path “Parent / Child” when the name is not unique. position: name of an existing position. status: new (empty = new), active, on_leave.",
    "Dates: YYYY-MM-DD or DD.MM.YYYY. Phone: +998 71 000 00 00. Formulas, macros and links are rejected. At most 500 rows and 2 MB.",
    "The import only ADDS new employees; it never changes or merges existing ones. A matching full name or e-mail is a “possible duplicate” and is skipped by default.",
    "Departments and positions must already exist in the portal. The example rows are fictional: delete them and enter your own data.",
]
